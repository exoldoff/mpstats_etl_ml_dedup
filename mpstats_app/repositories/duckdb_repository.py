from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import datetime, timedelta
import io
import json
import logging
from pathlib import Path
import re
from threading import RLock
import time
from typing import Any, Literal
from uuid import uuid4

from duckdb import sqltypes as duckdb_sqltypes
import pandas as pd

from pipeline.repositories.sql_repository import (
    apply_migrations,
    connect,
    duckdb_transaction,
    measure_duckdb_operation,
    quote_identifier,
    resolve_duckdb_temp_directory,
    sql_literal,
    table_exists,
)
from pipeline.services.sales_filter_service import DEFAULT_SALES_MIN_QUANTILE, DEFAULT_SALES_MIN_UNITS

from mpstats_app.config import AppSettings
from mpstats_app.utils import clean_record, clean_records, quote_duckdb_name


SEARCH_COLUMNS = ("SKU", "Артикул", "Название", "Бренд", "Категория")
EXPORT_METADATA_COLUMNS = (
    "__project_name",
    "__year",
    "__month",
    "__marketplace_code",
    "__source_type",
    "__category_key",
    "__row_hash",
    "__business_row_hash",
)
TEXT_DB_TYPES = ("CHAR", "STRING", "TEXT", "VARCHAR")
CUBE_SALES_FILTER_COLUMNS = ("Продажи, шт", "Продажи", "sales")
CUBE_VOLUME_FILTER_COLUMNS = ("Объем, кг", "Объём, кг", "Объем, т", "Объём, т", "Объем", "Объём", "volume_kg", "volume_t", "volume")
IMPORT_METADATA_COLUMNS = (
    "__run_id",
    "__source_file",
    "__imported_at",
    "__project_name",
    "__year",
    "__month",
    "__marketplace_code",
    "__source_type",
    "__category_key",
    "__row_hash",
    "__business_row_hash",
)
HEAVY_SLICE_ROWS_LIMIT = 250_000
HEAVY_CATEGORY_ROWS_LIMIT = 1_000_000
REPORT_REVENUE_COLUMNS = ("Выручка, руб", "Выручка", "revenue")
REPORT_VOLUME_KG_COLUMNS = ("Объем, кг", "Объём, кг", "volume_kg")
REPORT_VOLUME_T_COLUMNS = ("Объем, т", "Объём, т", "volume_t")
REPORT_CLASSIFICATION_COLUMNS = ("Тип", "Подкатегория", "Вид", "Вид мяса", "Сегмент")
DEDUP_EXPORT_COLUMNS = (
    "ML-группа товара",
    "ML-группа фасовки",
    "ML-канонический SKU",
    "ML-dedup статус",
    "ML-dedup run",
)
DEDUP_EXPORT_FAMILY_TITLE_FUNCTION = "mpstats_dedup_family_title"
DEDUP_FAMILY_MIN_CATEGORY_TOKENS = 2
DEDUP_FAMILY_UNIT_TOKENS = frozenset(
    {
        "г",
        "гр",
        "кг",
        "мл",
        "л",
        "шт",
        "штуки",
        "штук",
        "штука",
        "уп",
        "упак",
        "упаковка",
    }
)
DEDUP_FAMILY_GENERIC_TOKENS = frozenset(
    {
        "100",
        "aroy",
        "aroyd",
        "extra",
        "virgin",
        "extravirgin",
        "для",
        "без",
        "на",
        "и",
        "в",
        "с",
        "из",
        "по",
        "от",
        "до",
        "нерафинированное",
        "нерафинированный",
        "нерафинированная",
        "рафинированное",
        "рафинированный",
        "рафинированная",
        "пищевое",
        "пищевой",
        "пищевые",
        "растительное",
        "растительный",
        "растительная",
        "холодного",
        "холодный",
        "холодная",
        "отжима",
        "отжим",
        "первого",
        "первый",
        "натуральное",
        "натуральный",
        "натуральная",
        "органическое",
        "органический",
        "органическая",
        "универсальное",
        "универсальный",
        "косметическое",
        "косметический",
        "тела",
        "волос",
        "еды",
        "жарки",
        "добавок",
        "индонезия",
        "таиланд",
        "арой",
        "аройд",
    }
)
DEDUP_PRODUCTS_COLUMNS = (
    "run_id",
    "project_name",
    "category_key",
    "category_name",
    "ml_family_id",
    "ml_pack_id",
    "row_level",
    "sort_order",
    "node_id",
    "canonical_node_id",
    "canonical_sku",
    "normalized_sku",
    "marketplace_code",
    "marketplace",
    "article",
    "sku",
    "brand",
    "subcategory",
    "unit_amount",
    "total_amount",
    "multipack_count",
    "sales_volume",
    "revenue",
    "source_row_count",
    "component_size",
    "ml_dedup_status",
    "confidence_score",
)
XLSX_MAX_DATA_ROWS_WITH_HEADER = 1_048_575
CSV_DECIMAL_COMMA_PROTECTED_COLUMNS = {
    "дата",
    "sku",
    "продавец",
    "категория",
    "бренд",
    "год",
    "месяц",
    "подкатегория",
    "тип",
}
CSV_DECIMAL_DOT_PATTERN = r"([0-9])\.([0-9])"
CSV_DECIMAL_COMMA_REPLACEMENT = r"\1,\2"
CUBE_SALES_MIN_QUANTILE = DEFAULT_SALES_MIN_QUANTILE
CUBE_SALES_MIN_UNITS = DEFAULT_SALES_MIN_UNITS
DEDUP_ELIGIBLE_CATEGORY_GROUPS = {
    "соус": ("sauces", "Соусы"),
    "соусы": ("sauces", "Соусы"),
    "кокосовое масло": ("coconut_oil", "Кокосовое масло"),
    "мыло": ("soap", "Мыло"),
}


class DuplicateCubeSliceError(ValueError):
    pass


RAW_EXPORT_DOUBLE_COLUMNS = (
    "Продажи, шт",
    "Продажи",
    "Выручка, руб",
    "Выручка",
    "Вес, кг",
    "Вес, кг (ед.)",
    "Вес, кг (сумм.)",
    "Объем, кг",
    "Объём, кг",
    "Объем, т",
    "Объём, т",
    "Объем",
    "Объём",
    "Средняя цена, руб",
    "Средняя цена",
    "Цена за кг",
    "Цена",
    "price",
    "revenue",
    "volume",
    "volume_kg",
    "volume_t",
    "sales",
)
RAW_EXPORT_INTEGER_COLUMNS = (
    "Количество магазинов",
    "Кол-во магазинов",
    "stores_count",
)
RAW_EXPORT_DATE_COLUMNS = (
    "Дата",
    "date",
    "data_actual_until",
)
RAW_EXPORT_TIMESTAMP_COLUMNS = (
    "Дата и время",
    "datetime",
    "timestamp",
)
LOGGER = logging.getLogger(__name__)


@dataclass(frozen=True)
class ExportResult:
    output_path: Path
    file_size_bytes: int
    duration_seconds: float
    row_count: int | None
    status: str
    format: str = "csv"
    error: str | None = None


def _table_column_types(con: Any, table_name: str) -> dict[str, str]:
    return {
        str(row[0]): str(row[1]).upper()
        for row in con.execute(
            """
            SELECT column_name, data_type
            FROM information_schema.columns
            WHERE table_schema = 'main' AND table_name = ?
            """,
            [table_name],
        ).fetchall()
    }


def _db_type_accepts_text(data_type: str) -> bool:
    upper_type = data_type.upper()
    return any(marker in upper_type for marker in TEXT_DB_TYPES)


def _first_existing_column(columns: list[str], candidates: tuple[str, ...]) -> str | None:
    existing = set(columns)
    for column in candidates:
        if column in existing:
            return column
    return None


def _stage_column_types(con: Any, table_name: str) -> dict[str, str]:
    return {
        str(row[0]): str(row[1]).upper()
        for row in con.execute(f"DESCRIBE {quote_identifier(table_name)}").fetchall()
    }


def _csv_scan_sql(csv_path: Path) -> str:
    return (
        "read_csv("
        f"{sql_literal(str(csv_path))}, "
        "delim=';', header=true, all_varchar=true, null_padding=true, ignore_errors=false"
        ")"
    )


def _number_expr(column: str, *, table_alias: str | None = None) -> str:
    quoted = quote_duckdb_name(column)
    if table_alias:
        quoted = f"{table_alias}.{quoted}"
    nbsp = "\u00a0"
    return (
        "TRY_CAST("
        f"REPLACE(REPLACE(REPLACE(CAST({quoted} AS VARCHAR), '{nbsp}', ''), ' ', ''), ',', '.') "
        "AS DOUBLE)"
    )


def _normalized_text_expr(column: str, *, table_alias: str | None = None) -> str:
    quoted = quote_duckdb_name(column)
    if table_alias:
        quoted = f"{table_alias}.{quoted}"
    return f"lower(trim(CAST({quoted} AS VARCHAR)))"


def _stage_single_text_value(con: Any, *, stage_table: str, column: str) -> str | None:
    text_expr = f"NULLIF(TRIM(CAST({quote_duckdb_name(column)} AS VARCHAR)), '')"
    row = con.execute(
        f"""
        SELECT MIN({text_expr}) AS value, COUNT(DISTINCT {_normalized_text_expr(column)}) AS variants
        FROM {quote_identifier(stage_table)}
        WHERE {text_expr} IS NOT NULL
        """
    ).fetchone()
    if not row or int(row[1] or 0) != 1:
        return None
    return str(row[0])


def _positive_import_filter(columns: list[str]) -> str:
    filters: list[str] = []
    sales_column = _first_existing_column(columns, CUBE_SALES_FILTER_COLUMNS)
    volume_column = _first_existing_column(columns, CUBE_VOLUME_FILTER_COLUMNS)
    if sales_column:
        filters.append(f"{_number_expr(sales_column)} > 0")
    if volume_column:
        filters.append(f"{_number_expr(volume_column)} > 0")
    return " AND ".join(filters) if filters else "TRUE"


def _sales_quantile_source_sql(raw_table: str, columns: list[str], positive_filter: str) -> str:
    sales_column = _first_existing_column(columns, CUBE_SALES_FILTER_COLUMNS)
    quoted_raw_table = quote_identifier(raw_table)
    if not sales_column:
        return f"SELECT * FROM {quoted_raw_table} WHERE {positive_filter}"

    sales_expr = _number_expr(sales_column)
    sales_min_units = float(CUBE_SALES_MIN_UNITS)
    if CUBE_SALES_MIN_QUANTILE is None:
        return f"""
            SELECT *
            FROM {quoted_raw_table}
            WHERE {positive_filter}
              AND {sales_expr} >= {sales_min_units}
        """

    return f"""
        SELECT * EXCLUDE (__sales_quantile_threshold)
        FROM (
            SELECT
                *,
                QUANTILE_CONT({sales_expr}, {float(CUBE_SALES_MIN_QUANTILE)}) OVER () AS __sales_quantile_threshold
            FROM {quoted_raw_table}
            WHERE {positive_filter}
        )
        WHERE {sales_expr} >= GREATEST(__sales_quantile_threshold, {sales_min_units})
    """


def _hash_expr(
    columns: list[str],
    *,
    project_name: str | None = None,
    year: int | None = None,
    month: int | None = None,
    marketplace_code: str | None = None,
    category_key: str | None = None,
    table_alias: str | None = None,
) -> str:
    parts: list[str] = []
    for value in (project_name, year, month, marketplace_code, category_key):
        if value is not None:
            parts.append(sql_literal(str(value)))
    parts.extend(
        f"COALESCE(CAST({table_alias + '.' if table_alias else ''}{quote_duckdb_name(column)} AS VARCHAR), '<NULL>')"
        for column in columns
    )
    return "sha1(concat_ws('|', " + ", ".join(parts) + "))"


def _weight_column_aliases(source_columns: list[str]) -> dict[str, str]:
    if "Вес, кг (сумм.)" not in source_columns:
        return {}
    aliases = {"Вес, кг (сумм.)": "Вес, кг"}
    if "Вес, кг" in source_columns and "Вес, кг (ед.)" not in source_columns:
        aliases["Вес, кг"] = "Вес, кг (ед.)"
    return aliases


def _stage_has_text_values(con: Any, *, stage_table: str, column: str) -> bool:
    text_expr = f"NULLIF(TRIM(CAST({quote_duckdb_name(column)} AS VARCHAR)), '')"
    numeric_expr = _number_expr(column)
    row = con.execute(
        f"""
        SELECT COUNT(*)
        FROM {quote_identifier(stage_table)}
        WHERE {text_expr} IS NOT NULL
          AND {numeric_expr} IS NULL
        """
    ).fetchone()
    return bool(row and int(row[0]) > 0)


def _cast_for_target_type(column: str, data_type: str, *, table_alias: str | None = None) -> str:
    quoted = quote_duckdb_name(column)
    if table_alias:
        quoted = f"{table_alias}.{quoted}"
    upper = data_type.upper()
    if _db_type_accepts_text(upper):
        return quoted
    if any(marker in upper for marker in ("INT", "DOUBLE", "FLOAT", "REAL", "DECIMAL", "NUMERIC")):
        return f"TRY_CAST({_number_expr(column, table_alias=table_alias)} AS {upper})"
    if "DATE" in upper or "TIME" in upper:
        return f"TRY_CAST({quoted} AS {upper})"
    if "BOOL" in upper:
        return f"TRY_CAST({quoted} AS BOOLEAN)"
    return quoted


def _create_products_stage(
    con: Any,
    *,
    csv_path: Path,
    stage_table: str,
    run_id: str,
    source_file: str,
    project_name: str | None = None,
    year: int | None = None,
    month: int | None = None,
    marketplace_code: str | None = None,
    source_type: str | None = None,
    category_key: str | None = None,
) -> list[str]:
    raw_table = f"{stage_table}_raw"
    con.execute(f"CREATE OR REPLACE TEMP TABLE {quote_identifier(raw_table)} AS SELECT * FROM {_csv_scan_sql(csv_path)}")
    raw_columns = list(_stage_column_types(con, raw_table))
    source_columns = [column for column in raw_columns if column not in IMPORT_METADATA_COLUMNS]
    if not source_columns:
        raise ValueError(f"В файле нет колонок для загрузки в DuckDB: {csv_path}")

    aliases = _weight_column_aliases(source_columns)
    select_parts = [
        f"{quote_duckdb_name(column)} AS {quote_duckdb_name(aliases[column])}" if column in aliases else quote_duckdb_name(column)
        for column in source_columns
    ]
    select_parts.extend(
        [
            f"{sql_literal(run_id)} AS {quote_duckdb_name('__run_id')}",
            f"{sql_literal(source_file)} AS {quote_duckdb_name('__source_file')}",
            f"now() AS {quote_duckdb_name('__imported_at')}",
        ]
    )
    if project_name is not None:
        select_parts.append(f"{sql_literal(project_name)} AS {quote_duckdb_name('__project_name')}")
    if year is not None:
        select_parts.append(f"{int(year)} AS {quote_duckdb_name('__year')}")
    if month is not None:
        select_parts.append(f"{int(month)} AS {quote_duckdb_name('__month')}")
    if marketplace_code is not None:
        select_parts.append(f"{sql_literal(marketplace_code)} AS {quote_duckdb_name('__marketplace_code')}")
    if source_type is not None:
        select_parts.append(f"{sql_literal(source_type)} AS {quote_duckdb_name('__source_type')}")
    if category_key is not None:
        select_parts.append(f"{sql_literal(category_key)} AS {quote_duckdb_name('__category_key')}")
    select_parts.append(
        f"{_hash_expr(source_columns, project_name=project_name, year=year, month=month, marketplace_code=marketplace_code, category_key=category_key)} "
        f"AS {quote_duckdb_name('__row_hash')}"
    )
    select_parts.append(
        f"{_hash_expr(source_columns, project_name=project_name, year=year, month=month, marketplace_code=marketplace_code)} "
        f"AS {quote_duckdb_name('__business_row_hash')}"
    )
    positive_filter = _positive_import_filter(source_columns)
    stage_source_sql = _sales_quantile_source_sql(raw_table, source_columns, positive_filter)
    con.execute(
        f"""
        CREATE OR REPLACE TEMP TABLE {quote_identifier(stage_table)} AS
        SELECT * EXCLUDE (__row_number)
        FROM (
            SELECT
                {", ".join(select_parts)},
                ROW_NUMBER() OVER (
                    PARTITION BY {_hash_expr(source_columns, project_name=project_name, year=year, month=month, marketplace_code=marketplace_code, category_key=category_key)}
                    ORDER BY {_hash_expr(source_columns, project_name=project_name, year=year, month=month, marketplace_code=marketplace_code, category_key=category_key)}
                ) AS __row_number
            FROM ({stage_source_sql}) AS filtered_raw
        )
        WHERE __row_number = 1
        """
    )
    return list(_stage_column_types(con, stage_table))


def _ensure_table_accepts_stage_columns(con: Any, *, table_name: str, quoted_table: str, stage_table: str, stage_columns: list[str]) -> dict[str, str]:
    column_types = _table_column_types(con, table_name)
    stage_types = _stage_column_types(con, stage_table)
    for column in stage_columns:
        column_name = str(column)
        if column_name not in column_types:
            con.execute(f"ALTER TABLE {quoted_table} ADD COLUMN {quote_duckdb_name(column_name)} {stage_types[column_name]}")
            column_types[column_name] = stage_types[column_name]
            continue
        if _db_type_accepts_text(column_types[column_name]):
            continue
        if column_name.startswith("__"):
            continue
        if _stage_has_text_values(con, stage_table=stage_table, column=column_name):
            con.execute(f"ALTER TABLE {quoted_table} ALTER COLUMN {quote_duckdb_name(column_name)} TYPE VARCHAR")
            column_types[column_name] = "VARCHAR"
    return column_types


def _business_deduplicate_condition(*, quoted_table: str, stage_alias: str = "s") -> str:
    return f"""
        NOT EXISTS (
            SELECT 1
            FROM {quoted_table} t
            WHERE t.{quote_duckdb_name('__project_name')} = {stage_alias}.{quote_duckdb_name('__project_name')}
              AND CAST(t.{quote_duckdb_name('__year')} AS INTEGER) = CAST({stage_alias}.{quote_duckdb_name('__year')} AS INTEGER)
              AND CAST(t.{quote_duckdb_name('__month')} AS INTEGER) = CAST({stage_alias}.{quote_duckdb_name('__month')} AS INTEGER)
              AND t.{quote_duckdb_name('__marketplace_code')} = {stage_alias}.{quote_duckdb_name('__marketplace_code')}
              AND t.{quote_duckdb_name('__business_row_hash')} = {stage_alias}.{quote_duckdb_name('__business_row_hash')}
        )
    """


def _business_columns(stage_columns: list[str]) -> list[str]:
    metadata = set(IMPORT_METADATA_COLUMNS)
    return [column for column in stage_columns if column not in metadata]


def _backfill_business_row_hash(
    con: Any,
    *,
    quoted_table: str,
    stage_columns: list[str],
    project_name: str,
    year: int,
    month: int,
    marketplace_code: str,
) -> None:
    business_columns = _business_columns(stage_columns)
    if not business_columns:
        return
    con.execute(
        f"""
        UPDATE {quoted_table} AS t
        SET {quote_duckdb_name('__business_row_hash')} = {_hash_expr(
            business_columns,
            project_name=project_name,
            year=year,
            month=month,
            marketplace_code=marketplace_code,
            table_alias='t',
        )}
        WHERE t.{quote_duckdb_name('__project_name')} = ?
          AND CAST(t.{quote_duckdb_name('__year')} AS INTEGER) = ?
          AND CAST(t.{quote_duckdb_name('__month')} AS INTEGER) = ?
          AND t.{quote_duckdb_name('__marketplace_code')} = ?
          AND (
              t.{quote_duckdb_name('__business_row_hash')} IS NULL
              OR TRIM(CAST(t.{quote_duckdb_name('__business_row_hash')} AS VARCHAR)) = ''
          )
        """,
        [project_name, int(year), int(month), marketplace_code],
    )


def _fetch_cube_entry_by_natural_key(
    con: Any,
    *,
    project_name: str,
    year: int,
    month: int,
    marketplace_code: str,
    category_name: str,
) -> dict[str, Any] | None:
    result = con.execute(
        """
        SELECT *
        FROM cube_registry
        WHERE project_name = ?
          AND year = ?
          AND month = ?
          AND marketplace_code = ?
          AND lower(trim(CAST(category_name AS VARCHAR))) = lower(trim(CAST(? AS VARCHAR)))
        ORDER BY saved_to_db_at DESC
        LIMIT 1
        """,
        [project_name, int(year), int(month), marketplace_code, category_name],
    )
    row = result.fetchone()
    if not row:
        return None
    columns = [column[0] for column in result.description]
    return clean_record(dict(zip(columns, row)))


def _product_slice_count_by_natural_key(
    con: Any,
    *,
    table_name: str,
    quoted_table: str,
    project_name: str,
    year: int,
    month: int,
    marketplace_code: str,
    category_name: str,
) -> int:
    columns = set(_table_column_types(con, table_name))
    required = {"__project_name", "__year", "__month", "__marketplace_code", "Категория"}
    if not required.issubset(columns):
        return 0
    row = con.execute(
        f"""
        SELECT COUNT(*)
        FROM {quoted_table}
        WHERE {quote_duckdb_name('__project_name')} = ?
          AND CAST({quote_duckdb_name('__year')} AS INTEGER) = ?
          AND CAST({quote_duckdb_name('__month')} AS INTEGER) = ?
          AND {quote_duckdb_name('__marketplace_code')} = ?
          AND {_normalized_text_expr('Категория')} = lower(trim(CAST(? AS VARCHAR)))
        """,
        [project_name, int(year), int(month), marketplace_code, category_name],
    ).fetchone()
    return int(row[0] or 0) if row else 0


def _delete_product_slice_by_natural_key(
    con: Any,
    *,
    table_name: str,
    quoted_table: str,
    project_name: str,
    year: int,
    month: int,
    marketplace_code: str,
    category_name: str,
) -> int:
    rows = _product_slice_count_by_natural_key(
        con,
        table_name=table_name,
        quoted_table=quoted_table,
        project_name=project_name,
        year=year,
        month=month,
        marketplace_code=marketplace_code,
        category_name=category_name,
    )
    if not rows:
        return 0
    con.execute(
        f"""
        DELETE FROM {quoted_table}
        WHERE {quote_duckdb_name('__project_name')} = ?
          AND CAST({quote_duckdb_name('__year')} AS INTEGER) = ?
          AND CAST({quote_duckdb_name('__month')} AS INTEGER) = ?
          AND {quote_duckdb_name('__marketplace_code')} = ?
          AND {_normalized_text_expr('Категория')} = lower(trim(CAST(? AS VARCHAR)))
        """,
        [project_name, int(year), int(month), marketplace_code, category_name],
    )
    return rows


def _duplicate_cube_slice_message(*, project_name: str, year: int, month: int, marketplace_code: str, category_name: str) -> str:
    return (
        "Срез уже сохранён в кубе: "
        f"проект={project_name}, месяц={year}-{month:02d}, "
        f"маркетплейс={marketplace_code}, категория={category_name}. "
        "Удалите старый срез в разделе Данные -> Куб или включите Перезаписывать БД для осознанной замены."
    )


def _insert_stage_sql(
    *,
    quoted_table: str,
    stage_table: str,
    columns: list[str],
    target_types: dict[str, str],
    deduplicate: bool,
    business_deduplicate: bool = False,
) -> str:
    quoted_columns = ", ".join(quote_duckdb_name(column) for column in columns)
    selected_columns = ", ".join(f"{_cast_for_target_type(column, target_types[column], table_alias='s')} AS {quote_duckdb_name(column)}" for column in columns)
    conditions: list[str] = []
    if deduplicate:
        conditions.append(
            f"""NOT EXISTS (
            SELECT 1
            FROM {quoted_table} t
            WHERE t.{quote_duckdb_name('__row_hash')} = s.{quote_duckdb_name('__row_hash')}
        )"""
        )
    if business_deduplicate:
        conditions.append(_business_deduplicate_condition(quoted_table=quoted_table).strip())
    dedupe_sql = f"WHERE {' AND '.join(conditions)}" if conditions else ""
    return f"""
        INSERT INTO {quoted_table} ({quoted_columns})
        SELECT {selected_columns}
        FROM {quote_identifier(stage_table)} s
        {dedupe_sql}
    """


def _stage_count(con: Any, stage_table: str) -> int:
    row = con.execute(f"SELECT COUNT(*) FROM {quote_identifier(stage_table)}").fetchone()
    return int(row[0]) if row else 0


def _stage_insertable_count(con: Any, *, quoted_table: str, stage_table: str, business_deduplicate: bool) -> int:
    if not business_deduplicate:
        return _stage_count(con, stage_table)
    row = con.execute(
        f"""
        SELECT COUNT(*)
        FROM {quote_identifier(stage_table)} s
        WHERE {_business_deduplicate_condition(quoted_table=quoted_table).strip()}
        """
    ).fetchone()
    return int(row[0]) if row else 0


def _csv_header_prefix(columns: list[str], *, delimiter: str = ";") -> str:
    buffer = io.StringIO()
    csv.writer(buffer, delimiter=delimiter, lineterminator="\n").writerow(columns)
    return "\ufeff" + buffer.getvalue()


def _clean_copy_query(query: str) -> str:
    cleaned = query.strip()
    if cleaned.endswith(";"):
        cleaned = cleaned[:-1].rstrip()
    if not cleaned:
        raise ValueError("SQL-запрос для CSV-экспорта пуст.")
    return cleaned


def _csv_decimal_column_key(column: str) -> str:
    return str(column).strip().casefold()


def _csv_decimal_comma_expr(column: str) -> str:
    quoted = quote_duckdb_name(column)
    if _csv_decimal_column_key(column) in CSV_DECIMAL_COMMA_PROTECTED_COLUMNS:
        return f"{quoted} AS {quoted}"
    return (
        "CASE "
        f"WHEN {quoted} IS NULL THEN NULL "
        f"ELSE regexp_replace(CAST({quoted} AS VARCHAR), {sql_literal(CSV_DECIMAL_DOT_PATTERN)}, {sql_literal(CSV_DECIMAL_COMMA_REPLACEMENT)}, 'g') "
        f"END AS {quoted}"
    )


def _csv_decimal_comma_query(query: str, columns: list[str]) -> str:
    select_sql = ", ".join(_csv_decimal_comma_expr(column) for column in columns)
    return f"SELECT {select_sql} FROM ({query}) AS csv_decimal_source"


def _copy_row_count(rows: list[Any]) -> int | None:
    if len(rows) != 1:
        return None
    row = rows[0]
    if not row:
        return None
    try:
        return int(row[0])
    except (TypeError, ValueError):
        return None


def _copy_first_row_count(row: Any) -> int | None:
    if not row:
        return None
    try:
        return int(row[0])
    except (TypeError, ValueError):
        return None


def _normalize_query_params(params: dict[str, Any] | list[Any] | tuple[Any, ...] | None) -> dict[str, Any] | list[Any]:
    if params is None:
        return []
    if isinstance(params, dict):
        return dict(params)
    return list(params)


def _validate_csv_delimiter(delimiter: str) -> str:
    clean = str(delimiter or ";")
    if len(clean) != 1 or clean in {"\n", "\r"}:
        raise ValueError("CSV-разделитель должен быть одним символом без перевода строки.")
    return clean


def _validate_sheet_name(sheet_name: str) -> str:
    clean = str(sheet_name or "Data").strip() or "Data"
    for forbidden in ("\\", "/", "?", "*", "[", "]", ":"):
        clean = clean.replace(forbidden, "_")
    return clean[:31] or "Data"


def _clean_dedup_level(level: str) -> str:
    clean = str(level or "expanded").strip().casefold()
    if clean in {"canonical", "family"}:
        return clean
    return "expanded"


_DEDUP_FAMILY_TOKEN_RE = re.compile(r"[0-9a-zа-я]+")


def _dedup_family_tokens(value: object, *, keep_short: bool = False) -> list[str]:
    text = str(value or "").casefold().replace("ё", "е")
    tokens = _DEDUP_FAMILY_TOKEN_RE.findall(text)
    result: list[str] = []
    for token in tokens:
        if not keep_short and len(token) < 2:
            continue
        if token.isdigit():
            continue
        if token in DEDUP_FAMILY_UNIT_TOKENS:
            continue
        result.append(token)
    return result


def _dedup_brand_alias_tokens(brand: object) -> set[str]:
    tokens = set(_dedup_family_tokens(brand))
    normalized = " ".join(_dedup_family_tokens(brand, keep_short=True))
    if normalized in {"aroy d", "aroyd"} or "aroy" in tokens:
        tokens.update({"арой", "аройд"})
    return tokens


def _dedup_brand_pattern(brand: object) -> str | None:
    tokens = _dedup_family_tokens(brand, keep_short=True)
    if not tokens:
        return None
    return r"(?i)(?:[,;:\-–—]\s*)?\b" + r"[^0-9a-zа-я]*".join(re.escape(token) for token in tokens) + r"\b"


def _cleanup_dedup_family_title(value: str) -> str:
    title = re.sub(r"\s+([,;:])", r"\1", value)
    title = re.sub(r"\s{2,}", " ", title)
    title = re.sub(r"^[,;:\-–—]\s*", "", title)
    title = re.sub(r"[,;:\-–—]\s*$", "", title)
    return title.strip()


def _dedup_family_title(value: object, brand: object = None, category_name: object = None) -> str:
    title = str(value or "").strip()
    if not title:
        return ""
    original_title = title
    title = re.sub(
        r"(?i)(?:[,;:\-–—]\s*)?\bнабор\s*:?\s*\d+\s*(?:штук[аи]?|шт\.?)\b",
        "",
        title,
    )
    title = re.sub(r"(?i)(?:[,;:\-–—]\s*)?\b\d+\s*(?:штук[аи]?|шт\.?)\b", "", title)
    title = re.sub(r"(?i)(?:[,;:\-–—]\s*)?\b\d+(?:[,.]\d+)?\s*(?:мл|л|г|гр|кг)\b", "", title)
    title = re.sub(r"(?i)(?:[,;:\-–—]\s*)?\b\d+(?:[,.]\d+)?\s*%\b", "", title)
    brand_pattern = _dedup_brand_pattern(brand)
    if brand_pattern:
        title = re.sub(brand_pattern, "", title)
    if "aroy" in _dedup_brand_alias_tokens(brand):
        title = re.sub(r"(?i)(?:[,;:\-–—]\s*)?\bарой\s*[-–—]?\s*д\b", "", title)
    title = _cleanup_dedup_family_title(title)

    category_title = str(category_name or "").strip()
    category_tokens = set(_dedup_family_tokens(category_title))
    title_tokens = set(_dedup_family_tokens(title or original_title))
    if len(category_tokens) >= DEDUP_FAMILY_MIN_CATEGORY_TOKENS and category_tokens.issubset(title_tokens):
        extra_tokens = title_tokens - category_tokens - _dedup_brand_alias_tokens(brand)
        extra_tokens = {token for token in extra_tokens if token not in DEDUP_FAMILY_GENERIC_TOKENS}
        if not extra_tokens:
            return category_title

    return title or original_title


def _register_dedup_export_functions(con: Any) -> None:
    try:
        con.create_function(
            DEDUP_EXPORT_FAMILY_TITLE_FUNCTION,
            _dedup_family_title,
            return_type=duckdb_sqltypes.VARCHAR,
            null_handling="special",
        )
    except Exception as exc:
        if "already exists" not in str(exc).casefold():
            raise


def _dedup_category_group(category_name: object) -> tuple[str, str] | None:
    clean = " ".join(str(category_name or "").strip().split()).casefold().replace("ё", "е")
    return DEDUP_ELIGIBLE_CATEGORY_GROUPS.get(clean)


def _dedup_group_category_key(group_key: str) -> str:
    return f"dedupcat_{group_key}"


def _max_optional(left: object, right: object) -> object:
    if left in (None, ""):
        return right
    if right in (None, ""):
        return left
    try:
        return max(left, right)
    except TypeError:
        return max(str(left), str(right))


def _query_params_with_limit(
    params: dict[str, Any] | list[Any],
    *,
    limit: int,
    offset: int,
) -> tuple[str, dict[str, Any] | list[Any]]:
    if isinstance(params, dict):
        next_params = dict(params)
        next_params["__duckdb_limit"] = int(limit)
        next_params["__duckdb_offset"] = int(offset)
        return " LIMIT $__duckdb_limit OFFSET $__duckdb_offset", next_params
    return " LIMIT ? OFFSET ?", [*params, int(limit), int(offset)]


def _xlsx_cell_value(value: Any) -> Any:
    if pd.isna(value):
        return None
    return value


def _raw_export_column_expr(column: str) -> str:
    quoted = quote_duckdb_name(column)
    alias = quote_duckdb_name(column)
    lower = column.casefold()
    if column in RAW_EXPORT_INTEGER_COLUMNS or lower in {item.casefold() for item in RAW_EXPORT_INTEGER_COLUMNS}:
        return f"TRY_CAST({_number_expr(column)} AS BIGINT) AS {alias}"
    if column in RAW_EXPORT_DOUBLE_COLUMNS or lower in {item.casefold() for item in RAW_EXPORT_DOUBLE_COLUMNS}:
        return f"TRY_CAST({_number_expr(column)} AS DOUBLE) AS {alias}"
    if column in RAW_EXPORT_TIMESTAMP_COLUMNS or lower in {item.casefold() for item in RAW_EXPORT_TIMESTAMP_COLUMNS}:
        return f"TRY_CAST(NULLIF(TRIM(CAST({quoted} AS VARCHAR)), '') AS TIMESTAMP) AS {alias}"
    if column in RAW_EXPORT_DATE_COLUMNS or lower in {item.casefold() for item in RAW_EXPORT_DATE_COLUMNS} or ("дата" in lower and "время" not in lower):
        return f"NULLIF(TRIM(CAST({quoted} AS VARCHAR)), '') AS {alias}"
    return quoted


def _period_index_to_label(index: int) -> str:
    year = (index - 1) // 12
    month = (index - 1) % 12 + 1
    return f"{year}-{month:02d}"


def _max_iso(left: Any, right: Any) -> Any:
    if left is None:
        return right
    if right is None:
        return left
    return max(str(left), str(right))


class DuckDbAppRepository:
    def __init__(self, settings: AppSettings) -> None:
        self.settings = settings
        self._lock = RLock()

    def _duckdb_temp_directory(self) -> Path | None:
        return resolve_duckdb_temp_directory(fallback_directory=self.settings.project_root / "data" / "duckdb_tmp")

    def ensure_ready(self) -> None:
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            self._migrate_weight_columns(con, self.settings.products_table)

    def _migrate_weight_columns(self, con: Any, table_name: str) -> None:
        if not table_exists(con, table_name):
            return
        columns = _table_column_types(con, table_name)
        if "Вес, кг" not in columns:
            return
        quoted_table = quote_identifier(table_name)
        if "Вес, кг (ед.)" not in columns:
            con.execute(f"ALTER TABLE {quoted_table} ADD COLUMN {quote_duckdb_name('Вес, кг (ед.)')} DOUBLE")
        if "Вес, кг (сумм.)" in columns:
            con.execute(
                f"""
                UPDATE {quoted_table}
                SET
                    {quote_duckdb_name('Вес, кг (ед.)')} = COALESCE({quote_duckdb_name('Вес, кг (ед.)')}, TRY_CAST({quote_duckdb_name('Вес, кг')} AS DOUBLE)),
                    {quote_duckdb_name('Вес, кг')} = COALESCE(TRY_CAST({quote_duckdb_name('Вес, кг (сумм.)')} AS DOUBLE), TRY_CAST({quote_duckdb_name('Вес, кг')} AS DOUBLE))
                WHERE {quote_duckdb_name('Вес, кг (сумм.)')} IS NOT NULL OR {quote_duckdb_name('Вес, кг (ед.)')} IS NULL
                """
            )
            con.execute(f"ALTER TABLE {quoted_table} DROP COLUMN {quote_duckdb_name('Вес, кг (сумм.)')}")

    def _fetch_records(
        self,
        query: str,
        params: list[Any] | None = None,
        *,
        read_only: bool = False,
        temp_directory: Path | None = None,
        register_dedup_export_functions: bool = False,
    ) -> list[dict[str, Any]]:
        with self._lock, connect(self.settings.db_path, read_only=read_only, temp_directory=temp_directory) as con:
            if not read_only:
                apply_migrations(con)
            if register_dedup_export_functions:
                _register_dedup_export_functions(con)
            result = con.execute(query, params or [])
            columns = [col[0] for col in result.description]
            return clean_records([dict(zip(columns, row)) for row in result.fetchall()])

    def _fetch_one(
        self,
        query: str,
        params: list[Any] | None = None,
        *,
        read_only: bool = False,
        temp_directory: Path | None = None,
        register_dedup_export_functions: bool = False,
    ) -> dict[str, Any] | None:
        rows = self._fetch_records(
            query,
            params,
            read_only=read_only,
            temp_directory=temp_directory,
            register_dedup_export_functions=register_dedup_export_functions,
        )
        return rows[0] if rows else None

    def create_run(
        self,
        *,
        run_id: str,
        project_name: str,
        steps: str,
        source: str,
        schedule_id: str | None,
        workdir: Path,
        config_path: Path,
        rules_path: Path,
        db_path: Path,
        products_table: str,
        write_xlsx: bool,
        max_weight_kg: float,
        fill_unclassified: dict[str, Any] | None,
    ) -> dict[str, Any]:
        payload = json.dumps(fill_unclassified, ensure_ascii=False) if fill_unclassified else None
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute(
                """
                INSERT INTO app_runs (
                    run_id, project_name, steps, status, source, schedule_id, workdir,
                    config_path, rules_path, db_path, products_table, write_xlsx,
                    max_weight_kg, fill_unclassified_json
                )
                VALUES (?, ?, ?, 'queued', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    run_id,
                    project_name,
                    steps,
                    source,
                    schedule_id,
                    str(workdir),
                    str(config_path),
                    str(rules_path),
                    str(db_path),
                    products_table,
                    write_xlsx,
                    max_weight_kg,
                    payload,
                ],
            )
        self.add_event(run_id, "info", "Прогон поставлен в очередь", {"steps": steps, "source": source})
        return self.get_run(run_id) or {}

    def list_runs(self, *, limit: int = 50) -> list[dict[str, Any]]:
        return self._fetch_records(
            """
            SELECT *
            FROM app_runs
            ORDER BY created_at DESC
            LIMIT ?
            """,
            [int(limit)],
        )

    def get_run(self, run_id: str) -> dict[str, Any] | None:
        return self._fetch_one("SELECT * FROM app_runs WHERE run_id = ?", [run_id])

    def set_setting(self, key: str, value: str | None) -> None:
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute(
                """
                INSERT INTO app_settings (key, value, updated_at)
                VALUES (?, ?, now())
                ON CONFLICT (key) DO UPDATE SET value = EXCLUDED.value, updated_at = now()
                """,
                [key, value],
            )

    def get_setting(self, key: str) -> str | None:
        row = self._fetch_one("SELECT value FROM app_settings WHERE key = ?", [key])
        return str(row["value"]) if row and row.get("value") is not None else None

    def create_dedup_run(self, payload: dict[str, Any]) -> dict[str, Any]:
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute(
                """
                INSERT INTO dedup_runs (
                    run_id, project_name, category_key, category_name, status,
                    model_method, model_path, hf_model_id, embedding_model_name,
                    activation, threshold_strategy, threshold_same, faiss_top_k,
                    manifest_json
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    payload["run_id"],
                    payload["project_name"],
                    payload["category_key"],
                    payload.get("category_name"),
                    payload.get("status", "queued"),
                    payload["model_method"],
                    payload.get("model_path"),
                    payload.get("hf_model_id"),
                    payload.get("embedding_model_name"),
                    payload.get("activation"),
                    payload["threshold_strategy"],
                    float(payload["threshold_same"]),
                    int(payload["faiss_top_k"]),
                    json.dumps(payload.get("manifest") or {}, ensure_ascii=False),
                ],
            )
        return self.get_dedup_run(str(payload["run_id"])) or {}

    def update_dedup_run(self, run_id: str, values: dict[str, Any]) -> dict[str, Any] | None:
        allowed = {
            "status",
            "node_count",
            "candidate_count",
            "edge_count",
            "group_count",
            "manifest_path",
            "manifest_json",
            "error",
            "started_at",
            "finished_at",
        }
        assignments = [key for key in values if key in allowed]
        if not assignments:
            return self.get_dedup_run(run_id)
        sql = ", ".join(f"{key} = ?" for key in assignments)
        params: list[Any] = []
        for key in assignments:
            value = values[key]
            if key == "manifest_json" and isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False)
            params.append(value)
        params.append(run_id)
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute(f"UPDATE dedup_runs SET {sql} WHERE run_id = ?", params)
        return self.get_dedup_run(run_id)

    def get_dedup_run(self, run_id: str) -> dict[str, Any] | None:
        return self._fetch_one("SELECT * FROM dedup_runs WHERE run_id = ?", [run_id])

    def list_dedup_runs(
        self,
        *,
        project_name: str | None = None,
        category_key: str | None = None,
        limit: int = 50,
    ) -> list[dict[str, Any]]:
        where: list[str] = []
        params: list[Any] = []
        if project_name:
            where.append("project_name = ?")
            params.append(project_name)
        if category_key:
            where.append("category_key = ?")
            params.append(category_key)
        where_sql = "WHERE " + " AND ".join(where) if where else ""
        params.append(max(1, min(int(limit), 500)))
        return self._fetch_records(
            f"""
            SELECT *
            FROM dedup_runs
            {where_sql}
            ORDER BY created_at DESC
            LIMIT ?
            """,
            params,
        )

    def delete_dedup_runs(
        self,
        *,
        project_name: str,
        category_key: str | None = None,
        category_keys: list[str] | None = None,
        statuses: list[str],
    ) -> int:
        clean_statuses = [str(status).strip() for status in statuses if str(status).strip()]
        clean_category_keys = [
            str(key).strip()
            for key in (category_keys if category_keys is not None else [category_key])
            if str(key or "").strip()
        ]
        if not clean_statuses:
            return 0
        if not clean_category_keys:
            return 0
        status_placeholders = ", ".join("?" for _ in clean_statuses)
        category_placeholders = ", ".join("?" for _ in clean_category_keys)
        params: list[Any] = [project_name, *clean_category_keys, *clean_statuses]
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            run_rows = con.execute(
                f"""
                SELECT run_id
                FROM dedup_runs
                WHERE project_name = ?
                  AND category_key IN ({category_placeholders})
                  AND status IN ({status_placeholders})
                """,
                params,
            ).fetchall()
            run_ids = [str(row[0]) for row in run_rows if row and row[0]]
            if not run_ids:
                return 0
            run_placeholders = ", ".join("?" for _ in run_ids)
            for table_name in (
                "mpstats_products_dedup",
                "dedup_sku_groups",
                "dedup_sku_edges",
                "dedup_sku_nodes",
                "dedup_runs",
            ):
                con.execute(
                    f"DELETE FROM {table_name} WHERE run_id IN ({run_placeholders})",
                    run_ids,
                )
        return len(run_ids)

    def prune_failed_dedup_runs(self) -> int:
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            run_rows = con.execute(
                """
                SELECT run_id
                FROM (
                    SELECT
                        run_id,
                        ROW_NUMBER() OVER (
                            PARTITION BY project_name, COALESCE(category_name, category_key)
                            ORDER BY COALESCE(finished_at, created_at) DESC, created_at DESC
                        ) AS rn
                    FROM dedup_runs
                    WHERE status = 'failed'
                )
                WHERE rn > 1
                """
            ).fetchall()
            run_ids = [str(row[0]) for row in run_rows if row and row[0]]
            if not run_ids:
                return 0
            run_placeholders = ", ".join("?" for _ in run_ids)
            for table_name in (
                "mpstats_products_dedup",
                "dedup_sku_groups",
                "dedup_sku_edges",
                "dedup_sku_nodes",
                "dedup_runs",
            ):
                con.execute(
                    f"DELETE FROM {table_name} WHERE run_id IN ({run_placeholders})",
                    run_ids,
                )
        return len(run_ids)

    def fail_stale_dedup_runs(self, *, project_name: str | None = None) -> int:
        where = ["status IN ('queued', 'running')"]
        params: list[Any] = []
        if project_name:
            where.append("project_name = ?")
            params.append(project_name)
        where_sql = " AND ".join(where)
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            count_row = con.execute(
                f"SELECT COUNT(*) AS runs_count FROM dedup_runs WHERE {where_sql}",
                params,
            ).fetchone()
            runs_count = int(count_row[0] or 0) if count_row else 0
            if runs_count <= 0:
                return 0
            con.execute(
                f"""
                UPDATE dedup_runs
                SET
                    status = 'failed',
                    error = COALESCE(error, 'Dedup run interrupted: backend restarted before completion.'),
                    finished_at = COALESCE(finished_at, now())
                WHERE {where_sql}
                """,
                params,
            )
        return runs_count

    def latest_successful_dedup_run(self, *, project_name: str, category_key: str) -> dict[str, Any] | None:
        return self._fetch_one(
            """
            SELECT *
            FROM dedup_runs
            WHERE project_name = ? AND category_key = ? AND status = 'success'
            ORDER BY COALESCE(finished_at, created_at) DESC, created_at DESC
            LIMIT 1
            """,
            [project_name, category_key],
        )

    def list_dedup_eligible_categories(self, *, project_name: str) -> list[dict[str, Any]]:
        rows = self._fetch_records(
            """
            SELECT
                category_key,
                MIN(category_name) AS category_name,
                MIN(marketplace_code) AS marketplace_code,
                MIN(marketplace) AS marketplace,
                SUM(rows_count) AS rows_count,
                COUNT(*) AS slices_count,
                MAX(saved_to_db_at) AS latest_saved_at,
                MAX(COALESCE(exported_at, saved_to_db_at)) AS latest_source_at
            FROM cube_registry
            WHERE project_name = ?
              AND lower(trim(CAST(category_name AS VARCHAR))) IN (
                'соус', 'соусы', 'кокосовое масло', 'мыло'
              )
            GROUP BY category_key
            ORDER BY category_name, category_key
            """,
            [project_name],
        )
        grouped: dict[str, dict[str, Any]] = {}
        source_sets: dict[str, set[str]] = {}
        marketplace_sets: dict[str, set[str]] = {}
        marketplace_code_sets: dict[str, set[str]] = {}
        for row in rows:
            group = _dedup_category_group(row.get("category_name"))
            if group is None:
                continue
            group_key, group_name = group
            item = grouped.setdefault(
                group_key,
                {
                    "category_key": _dedup_group_category_key(group_key),
                    "category_name": group_name,
                    "rows_count": 0,
                    "slices_count": 0,
                    "latest_saved_at": None,
                    "latest_source_at": None,
                },
            )
            item["rows_count"] = int(item["rows_count"] or 0) + int(row.get("rows_count") or 0)
            item["slices_count"] = int(item["slices_count"] or 0) + int(row.get("slices_count") or 0)
            item["latest_saved_at"] = _max_optional(item.get("latest_saved_at"), row.get("latest_saved_at"))
            item["latest_source_at"] = _max_optional(item.get("latest_source_at"), row.get("latest_source_at"))

            source_key = str(row.get("category_key") or "").strip()
            if source_key:
                source_sets.setdefault(group_key, set()).add(source_key)
            marketplace = str(row.get("marketplace") or "").strip()
            if marketplace:
                marketplace_sets.setdefault(group_key, set()).add(marketplace)
            marketplace_code = str(row.get("marketplace_code") or "").strip()
            if marketplace_code:
                marketplace_code_sets.setdefault(group_key, set()).add(marketplace_code)

        result: list[dict[str, Any]] = []
        for group_key, item in grouped.items():
            source_category_keys = sorted(source_sets.get(group_key, set()))
            marketplaces = sorted(marketplace_sets.get(group_key, set()))
            marketplace_codes = sorted(marketplace_code_sets.get(group_key, set()))
            result.append(
                {
                    **item,
                    "source_category_keys": source_category_keys,
                    "source_categories_count": len(source_category_keys),
                    "marketplaces": marketplaces,
                    "marketplace_codes": marketplace_codes,
                }
            )
        return sorted(result, key=lambda item: (str(item.get("category_name") or ""), str(item.get("category_key") or "")))

    def resolve_dedup_source_category_keys(self, *, project_name: str, category_key: str) -> list[str]:
        clean_key = str(category_key or "").strip()
        if not clean_key:
            return []
        for category in self.list_dedup_eligible_categories(project_name=project_name):
            if str(category.get("category_key") or "") == clean_key:
                source_keys = [str(key).strip() for key in category.get("source_category_keys") or []]
                return [key for key in source_keys if key]
        return [clean_key]

    def fetch_dedup_source_dataframe(
        self,
        *,
        table_name: str,
        project_name: str,
        category_key: str | None = None,
        category_keys: list[str] | None = None,
    ) -> pd.DataFrame:
        columns = self.table_columns(table_name)
        self._require_export_metadata(columns)
        quoted_table = quote_identifier(table_name)
        clean_category_keys = [
            str(key).strip()
            for key in (category_keys if category_keys is not None else [category_key])
            if str(key or "").strip()
        ]
        if not clean_category_keys:
            return pd.DataFrame()
        category_placeholders = ", ".join("?" for _ in clean_category_keys)

        def text_expr(column: str, fallback: str = "NULL") -> str:
            if column in columns:
                return f"NULLIF(TRIM(CAST({quote_duckdb_name(column)} AS VARCHAR)), '')"
            return fallback

        article_expr = text_expr("Артикул", "NULL")
        if "__business_row_hash" in columns:
            article_expr = f"COALESCE({article_expr}, NULLIF(TRIM(CAST({quote_duckdb_name('__business_row_hash')} AS VARCHAR)), ''))"
        if "__row_hash" in columns:
            article_expr = f"COALESCE({article_expr}, NULLIF(TRIM(CAST({quote_duckdb_name('__row_hash')} AS VARCHAR)), ''))"

        title_expr = f"COALESCE({text_expr('SKU')}, {text_expr('Название')}, {article_expr})"
        marketplace_expr = f"COALESCE({text_expr('Маркетплейс')}, {text_expr('__marketplace_code')})"
        category_expr = f"COALESCE({text_expr('Категория')}, {text_expr('__category_key')})"
        unit_expr = _number_expr("Вес, кг (ед.)") if "Вес, кг (ед.)" in columns else "CAST(NULL AS DOUBLE)"
        total_expr = _number_expr("Вес, кг") if "Вес, кг" in columns else "CAST(NULL AS DOUBLE)"
        sales_col = _first_existing_column(columns, CUBE_SALES_FILTER_COLUMNS)
        volume_col = _first_existing_column(columns, CUBE_VOLUME_FILTER_COLUMNS)
        revenue_col = _first_existing_column(columns, REPORT_REVENUE_COLUMNS)
        sales_expr = _number_expr(sales_col) if sales_col else "CAST(0 AS DOUBLE)"
        volume_expr = _number_expr(volume_col) if volume_col else None
        revenue_expr = _number_expr(revenue_col) if revenue_col else "CAST(0 AS DOUBLE)"
        source_filters = [
            f"{quote_duckdb_name('__project_name')} = ?",
            f"{quote_duckdb_name('__category_key')} IN ({category_placeholders})",
        ]
        if sales_col:
            source_filters.append(f"{sales_expr} >= {float(CUBE_SALES_MIN_UNITS)}")
        if volume_expr:
            source_filters.append(f"{volume_expr} > 0")
        source_where_sql = " AND ".join(source_filters)

        with self._lock, connect(self.settings.db_path, read_only=True, temp_directory=self._duckdb_temp_directory()) as con:
            return con.execute(
                f"""
                SELECT
                    CAST({quote_duckdb_name('__project_name')} AS VARCHAR) AS project_name,
                    CAST({quote_duckdb_name('__category_key')} AS VARCHAR) AS category_key,
                    {category_expr} AS category_name,
                    CAST({quote_duckdb_name('__marketplace_code')} AS VARCHAR) AS marketplace_code,
                    {marketplace_expr} AS marketplace,
                    {article_expr} AS article,
                    {title_expr} AS sku,
                    {text_expr('Бренд')} AS brand,
                    {text_expr('Подкатегория')} AS subcategory,
                    {unit_expr} AS unit_amount,
                    {total_expr} AS total_amount,
                    {sales_expr} AS sales_volume,
                    {revenue_expr} AS revenue,
                    {text_expr('__row_hash')} AS row_hash,
                    {text_expr('__business_row_hash')} AS business_row_hash
                FROM {quoted_table}
                WHERE {source_where_sql}
                """,
                [project_name, *clean_category_keys],
            ).fetchdf()

    def replace_dedup_nodes(self, run_id: str, rows: list[dict[str, Any]]) -> None:
        self._replace_dedup_rows(
            table_name="dedup_sku_nodes",
            run_id=run_id,
            columns=[
                "run_id",
                "node_id",
                "project_name",
                "category_key",
                "category_name",
                "marketplace_code",
                "marketplace",
                "article",
                "sku",
                "brand",
                "subcategory",
                "unit_amount",
                "total_amount",
                "multipack_count",
                "sales_volume",
                "revenue",
                "row_count",
                "source_row_hashes_json",
                "embedding_text",
            ],
            rows=rows,
        )

    def replace_dedup_edges(self, run_id: str, rows: list[dict[str, Any]]) -> None:
        self._replace_dedup_rows(
            table_name="dedup_sku_edges",
            run_id=run_id,
            columns=[
                "run_id",
                "edge_id",
                "node_id_a",
                "node_id_b",
                "score",
                "threshold_strategy",
                "threshold_same",
                "predicted_binary",
                "predicted_label",
                "candidate_rank",
                "candidate_source",
                "blocking_scope",
                "same_pack_signature",
            ],
            rows=rows,
        )

    def replace_dedup_groups(self, run_id: str, rows: list[dict[str, Any]]) -> None:
        self._replace_dedup_rows(
            table_name="dedup_sku_groups",
            run_id=run_id,
            columns=[
                "run_id",
                "node_id",
                "ml_family_id",
                "ml_pack_id",
                "canonical_node_id",
                "canonical_sku",
                "ml_dedup_status",
                "confidence_score",
                "component_size",
            ],
            rows=rows,
        )

    def fetch_dedup_identity_assignments(
        self,
        *,
        project_name: str,
        category_key: str,
        model_method: str,
        model_path: str,
        hf_model_id: str,
        embedding_model_name: str,
        faiss_top_k: int,
        activation: str,
        threshold_strategy: str,
        threshold_same: float,
        node_ids: list[str],
    ) -> list[dict[str, Any]]:
        clean_node_ids = sorted({str(node_id).strip() for node_id in node_ids if str(node_id).strip()})
        if not clean_node_ids:
            return []
        placeholders = ", ".join("?" for _ in clean_node_ids)
        return self._fetch_records(
            f"""
            SELECT
                project_name,
                category_key,
                node_id,
                model_method,
                model_path,
                hf_model_id,
                embedding_model_name,
                faiss_top_k,
                activation,
                threshold_strategy,
                threshold_same,
                ml_family_id,
                ml_pack_id,
                canonical_node_id,
                canonical_sku,
                ml_dedup_status,
                confidence_score,
                component_size,
                source_run_id,
                updated_at
            FROM dedup_identity_assignments
            WHERE project_name = ?
              AND category_key = ?
              AND model_method = ?
              AND model_path = ?
              AND hf_model_id = ?
              AND embedding_model_name = ?
              AND faiss_top_k = ?
              AND activation = ?
              AND threshold_strategy = ?
              AND threshold_same = ?
              AND node_id IN ({placeholders})
            """,
            [
                project_name,
                category_key,
                model_method,
                model_path,
                hf_model_id,
                embedding_model_name,
                int(faiss_top_k),
                activation,
                threshold_strategy,
                float(threshold_same),
                *clean_node_ids,
            ],
            read_only=True,
            temp_directory=self._duckdb_temp_directory(),
        )

    def upsert_dedup_identity_assignments(
        self,
        *,
        project_name: str,
        category_key: str,
        model_method: str,
        model_path: str,
        hf_model_id: str,
        embedding_model_name: str,
        faiss_top_k: int,
        activation: str,
        threshold_strategy: str,
        threshold_same: float,
        run_id: str,
        rows: list[dict[str, Any]],
    ) -> None:
        if not rows:
            return
        columns = [
            "project_name",
            "category_key",
            "node_id",
            "model_method",
            "model_path",
            "hf_model_id",
            "embedding_model_name",
            "faiss_top_k",
            "activation",
            "threshold_strategy",
            "threshold_same",
            "ml_family_id",
            "ml_pack_id",
            "canonical_node_id",
            "canonical_sku",
            "ml_dedup_status",
            "confidence_score",
            "component_size",
            "source_run_id",
        ]
        values = [
            [
                project_name,
                category_key,
                row.get("node_id"),
                model_method,
                model_path,
                hf_model_id,
                embedding_model_name,
                int(faiss_top_k),
                activation,
                threshold_strategy,
                float(threshold_same),
                row.get("ml_family_id"),
                row.get("ml_pack_id"),
                row.get("canonical_node_id"),
                row.get("canonical_sku"),
                row.get("ml_dedup_status"),
                row.get("confidence_score"),
                row.get("component_size"),
                run_id,
            ]
            for row in rows
        ]
        placeholders = ", ".join("?" for _ in columns)
        column_sql = ", ".join(columns)
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.executemany(
                f"""
                INSERT INTO dedup_identity_assignments ({column_sql})
                VALUES ({placeholders})
                ON CONFLICT (project_name, category_key, node_id) DO UPDATE SET
                    model_method = EXCLUDED.model_method,
                    model_path = EXCLUDED.model_path,
                    hf_model_id = EXCLUDED.hf_model_id,
                    embedding_model_name = EXCLUDED.embedding_model_name,
                    faiss_top_k = EXCLUDED.faiss_top_k,
                    activation = EXCLUDED.activation,
                    threshold_strategy = EXCLUDED.threshold_strategy,
                    threshold_same = EXCLUDED.threshold_same,
                    ml_family_id = EXCLUDED.ml_family_id,
                    ml_pack_id = EXCLUDED.ml_pack_id,
                    canonical_node_id = EXCLUDED.canonical_node_id,
                    canonical_sku = EXCLUDED.canonical_sku,
                    ml_dedup_status = EXCLUDED.ml_dedup_status,
                    confidence_score = EXCLUDED.confidence_score,
                    component_size = EXCLUDED.component_size,
                    source_run_id = EXCLUDED.source_run_id,
                    updated_at = now()
                """,
                values,
            )

    def fetch_dedup_node_embeddings(
        self,
        *,
        project_name: str,
        category_key: str,
        embedding_model_name: str,
        text_builder_version: str,
        normalize_embeddings: bool,
        entries: list[dict[str, Any]],
    ) -> dict[str, dict[str, Any]]:
        clean_entries = [
            {
                "node_id": str(entry.get("node_id") or "").strip(),
                "embedding_text_hash": str(entry.get("embedding_text_hash") or "").strip(),
            }
            for entry in entries
        ]
        clean_entries = [entry for entry in clean_entries if entry["node_id"] and entry["embedding_text_hash"]]
        if not clean_entries:
            return {}
        lookup = pd.DataFrame(clean_entries)
        with self._lock, connect(self.settings.db_path, temp_directory=self._duckdb_temp_directory()) as con:
            apply_migrations(con)
            con.register("_dedup_embedding_lookup", lookup)
            try:
                rows = con.execute(
                    """
                    SELECT
                        e.node_id,
                        e.embedding_text_hash,
                        e.embedding_dimension,
                        e.embedding_blob
                    FROM dedup_node_embeddings AS e
                    JOIN _dedup_embedding_lookup AS l
                      ON l.node_id = e.node_id
                     AND l.embedding_text_hash = e.embedding_text_hash
                    WHERE e.project_name = ?
                      AND e.category_key = ?
                      AND e.embedding_model_name = ?
                      AND e.text_builder_version = ?
                      AND e.normalize_embeddings = ?
                    """,
                    [project_name, category_key, embedding_model_name, text_builder_version, bool(normalize_embeddings)],
                ).fetchall()
            finally:
                con.unregister("_dedup_embedding_lookup")
        return {
            str(row[0]): {
                "embedding_text_hash": str(row[1]),
                "embedding_dimension": int(row[2]),
                "embedding_blob": bytes(row[3]),
            }
            for row in rows
        }

    def upsert_dedup_node_embeddings(
        self,
        *,
        project_name: str,
        category_key: str,
        embedding_model_name: str,
        text_builder_version: str,
        normalize_embeddings: bool,
        rows: list[dict[str, Any]],
    ) -> None:
        if not rows:
            return
        columns = [
            "project_name",
            "category_key",
            "node_id",
            "embedding_model_name",
            "text_builder_version",
            "normalize_embeddings",
            "embedding_text_hash",
            "embedding_dimension",
            "embedding_blob",
        ]
        values = [
            [
                project_name,
                category_key,
                row.get("node_id"),
                embedding_model_name,
                text_builder_version,
                bool(normalize_embeddings),
                row.get("embedding_text_hash"),
                int(row.get("embedding_dimension") or 0),
                row.get("embedding_blob"),
            ]
            for row in rows
        ]
        placeholders = ", ".join("?" for _ in columns)
        column_sql = ", ".join(columns)
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.executemany(
                f"""
                INSERT INTO dedup_node_embeddings ({column_sql})
                VALUES ({placeholders})
                ON CONFLICT (
                    project_name,
                    category_key,
                    node_id,
                    embedding_model_name,
                    text_builder_version,
                    normalize_embeddings,
                    embedding_text_hash
                ) DO UPDATE SET
                    embedding_dimension = EXCLUDED.embedding_dimension,
                    embedding_blob = EXCLUDED.embedding_blob,
                    updated_at = now()
                """,
                values,
            )

    def fetch_dedup_pair_scores(
        self,
        *,
        project_name: str,
        category_key: str,
        model_method: str,
        model_path: str,
        hf_model_id: str,
        activation: str,
        entries: list[dict[str, Any]],
    ) -> dict[tuple[str, str], float]:
        clean_entries = [
            {
                "node_id_a": str(entry.get("node_id_a") or "").strip(),
                "node_id_b": str(entry.get("node_id_b") or "").strip(),
                "pair_text_hash": str(entry.get("pair_text_hash") or "").strip(),
            }
            for entry in entries
        ]
        clean_entries = [
            entry
            for entry in clean_entries
            if entry["node_id_a"] and entry["node_id_b"] and entry["pair_text_hash"]
        ]
        if not clean_entries:
            return {}
        lookup = pd.DataFrame(clean_entries)
        with self._lock, connect(self.settings.db_path, temp_directory=self._duckdb_temp_directory()) as con:
            apply_migrations(con)
            con.register("_dedup_pair_lookup", lookup)
            try:
                rows = con.execute(
                    """
                    SELECT
                        s.node_id_a,
                        s.node_id_b,
                        s.score
                    FROM dedup_pair_score_cache AS s
                    JOIN _dedup_pair_lookup AS l
                      ON l.node_id_a = s.node_id_a
                     AND l.node_id_b = s.node_id_b
                     AND l.pair_text_hash = s.pair_text_hash
                    WHERE s.project_name = ?
                      AND s.category_key = ?
                      AND s.model_method = ?
                      AND s.model_path = ?
                      AND s.hf_model_id = ?
                      AND s.activation = ?
                    """,
                    [project_name, category_key, model_method, model_path, hf_model_id, activation],
                ).fetchall()
            finally:
                con.unregister("_dedup_pair_lookup")
        return {(str(row[0]), str(row[1])): float(row[2]) for row in rows}

    def upsert_dedup_pair_scores(
        self,
        *,
        project_name: str,
        category_key: str,
        model_method: str,
        model_path: str,
        hf_model_id: str,
        activation: str,
        rows: list[dict[str, Any]],
    ) -> None:
        if not rows:
            return
        columns = [
            "project_name",
            "category_key",
            "node_id_a",
            "node_id_b",
            "model_method",
            "model_path",
            "hf_model_id",
            "activation",
            "pair_text_hash",
            "score",
        ]
        values = [
            [
                project_name,
                category_key,
                row.get("node_id_a"),
                row.get("node_id_b"),
                model_method,
                model_path,
                hf_model_id,
                activation,
                row.get("pair_text_hash"),
                float(row.get("score") or 0),
            ]
            for row in rows
        ]
        placeholders = ", ".join("?" for _ in columns)
        column_sql = ", ".join(columns)
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.executemany(
                f"""
                INSERT INTO dedup_pair_score_cache ({column_sql})
                VALUES ({placeholders})
                ON CONFLICT (
                    project_name,
                    category_key,
                    node_id_a,
                    node_id_b,
                    model_method,
                    model_path,
                    hf_model_id,
                    activation,
                    pair_text_hash
                ) DO UPDATE SET
                    score = EXCLUDED.score,
                    updated_at = now()
                """,
                values,
            )

    def _replace_dedup_rows(
        self,
        *,
        table_name: str,
        run_id: str,
        columns: list[str],
        rows: list[dict[str, Any]],
    ) -> None:
        quoted_table = quote_identifier(table_name)
        placeholders = ", ".join("?" for _ in columns)
        column_sql = ", ".join(columns)
        values = [[row.get(column) for column in columns] for row in rows]
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            with duckdb_transaction(con):
                con.execute(f"DELETE FROM {quoted_table} WHERE run_id = ?", [run_id])
                if values:
                    con.executemany(
                        f"INSERT INTO {quoted_table} ({column_sql}) VALUES ({placeholders})",
                        values,
                    )

    def fetch_dedup_artifact(self, *, run_id: str, artifact: str, limit: int = 1000) -> list[dict[str, Any]]:
        if artifact == "edges":
            return self._fetch_records(
                """
                SELECT *
                FROM dedup_sku_edges
                WHERE run_id = ?
                ORDER BY predicted_binary DESC, score DESC NULLS LAST, edge_id
                LIMIT ?
                """,
                [run_id, max(1, min(int(limit), 100_000))],
                read_only=True,
                temp_directory=self._duckdb_temp_directory(),
            )
        if artifact == "groups":
            return self._fetch_records(
                """
                SELECT
                    g.*,
                    n.category_key,
                    n.category_name,
                    n.marketplace,
                    n.marketplace_code,
                    n.article,
                    n.sku,
                    n.brand,
                    n.subcategory,
                    n.sales_volume,
                    n.revenue
                FROM dedup_sku_groups AS g
                LEFT JOIN dedup_sku_nodes AS n
                  ON n.run_id = g.run_id AND n.node_id = g.node_id
                WHERE g.run_id = ?
                ORDER BY g.ml_family_id, g.ml_pack_id, n.sku
                LIMIT ?
                """,
                [run_id, max(1, min(int(limit), 100_000))],
                read_only=True,
                temp_directory=self._duckdb_temp_directory(),
            )
        raise ValueError("artifact должен быть groups или edges.")

    def refresh_dedup_products_table(self, *, run_id: str) -> int:
        run = self.get_dedup_run(run_id)
        if not run:
            raise ValueError("Dedup run не найден.")
        project_name = str(run["project_name"])
        category_key = str(run["category_key"])
        category_name = str(run.get("category_name") or "") or None
        source_category_keys = self.resolve_dedup_source_category_keys(project_name=project_name, category_key=category_key)
        delete_category_keys = sorted({category_key, *source_category_keys})
        delete_placeholders = ", ".join("?" for _ in delete_category_keys)
        with self._lock, connect(self.settings.db_path, temp_directory=self._duckdb_temp_directory()) as con:
            apply_migrations(con)
            with duckdb_transaction(con):
                con.execute(
                    f"""
                    DELETE FROM mpstats_products_dedup
                    WHERE project_name = ? AND category_key IN ({delete_placeholders})
                    """,
                    [project_name, *delete_category_keys],
                )
                con.execute(
                    """
                    INSERT INTO mpstats_products_dedup (
                        run_id, project_name, category_key, category_name,
                        ml_family_id, ml_pack_id, row_level, sort_order,
                        node_id, canonical_node_id, canonical_sku, normalized_sku,
                        marketplace_code, marketplace, article, sku, brand, subcategory,
                        unit_amount, total_amount, multipack_count,
                        sales_volume, revenue, source_row_count, component_size,
                        ml_dedup_status, confidence_score
                    )
                    SELECT
                        g.run_id,
                        n.project_name,
                        ? AS category_key,
                        COALESCE(?, MIN(n.category_name)) AS category_name,
                        g.ml_family_id,
                        g.ml_pack_id,
                        'canonical' AS row_level,
                        0 AS sort_order,
                        g.canonical_node_id AS node_id,
                        g.canonical_node_id,
                        MAX(g.canonical_sku) AS canonical_sku,
                        MAX(g.canonical_sku) AS normalized_sku,
                        MIN(cn.marketplace_code) AS marketplace_code,
                        MIN(cn.marketplace) AS marketplace,
                        MIN(cn.article) AS article,
                        MIN(cn.sku) AS sku,
                        MIN(cn.brand) AS brand,
                        MIN(cn.subcategory) AS subcategory,
                        MIN(cn.unit_amount) AS unit_amount,
                        MIN(cn.total_amount) AS total_amount,
                        MIN(cn.multipack_count) AS multipack_count,
                        COALESCE(SUM(n.sales_volume), 0) AS sales_volume,
                        COALESCE(SUM(n.revenue), 0) AS revenue,
                        COALESCE(SUM(n.row_count), 0) AS source_row_count,
                        COUNT(*) AS component_size,
                        CASE WHEN COUNT(*) > 1 THEN 'canonical_group' ELSE 'canonical_singleton' END AS ml_dedup_status,
                        MAX(g.confidence_score) AS confidence_score
                    FROM dedup_sku_groups AS g
                    JOIN dedup_sku_nodes AS n
                      ON n.run_id = g.run_id AND n.node_id = g.node_id
                    LEFT JOIN dedup_sku_nodes AS cn
                      ON cn.run_id = g.run_id AND cn.node_id = g.canonical_node_id
                    WHERE g.run_id = ?
                    GROUP BY
                        g.run_id,
                        n.project_name,
                        g.ml_family_id,
                        g.ml_pack_id,
                        g.canonical_node_id
                    """,
                    [category_key, category_name, run_id],
                )
                con.execute(
                    """
                    INSERT INTO mpstats_products_dedup (
                        run_id, project_name, category_key, category_name,
                        ml_family_id, ml_pack_id, row_level, sort_order,
                        node_id, canonical_node_id, canonical_sku, normalized_sku,
                        marketplace_code, marketplace, article, sku, brand, subcategory,
                        unit_amount, total_amount, multipack_count,
                        sales_volume, revenue, source_row_count, component_size,
                        ml_dedup_status, confidence_score
                    )
                    SELECT
                        g.run_id,
                        n.project_name,
                        n.category_key,
                        n.category_name,
                        g.ml_family_id,
                        g.ml_pack_id,
                        'member' AS row_level,
                        1 AS sort_order,
                        n.node_id,
                        g.canonical_node_id,
                        g.canonical_sku,
                        g.canonical_sku AS normalized_sku,
                        n.marketplace_code,
                        n.marketplace,
                        n.article,
                        n.sku,
                        n.brand,
                        n.subcategory,
                        n.unit_amount,
                        n.total_amount,
                        n.multipack_count,
                        n.sales_volume,
                        n.revenue,
                        n.row_count AS source_row_count,
                        g.component_size,
                        g.ml_dedup_status,
                        g.confidence_score
                    FROM dedup_sku_groups AS g
                    JOIN dedup_sku_nodes AS n
                      ON n.run_id = g.run_id AND n.node_id = g.node_id
                    WHERE g.run_id = ?
                    """,
                    [run_id],
                )
                row = con.execute(
                    """
                    SELECT COUNT(*) AS rows_count
                    FROM mpstats_products_dedup
                    WHERE run_id = ?
                    """,
                    [run_id],
                ).fetchone()
        return int(row[0]) if row else 0

    def fetch_dedup_products(
        self,
        *,
        project_name: str,
        category_key: str | None = None,
        level: str = "expanded",
        query_text: str | None = None,
        limit: int = 500,
    ) -> dict[str, Any]:
        if not self.table_exists("mpstats_products_dedup"):
            return {"columns": list(DEDUP_PRODUCTS_COLUMNS), "rows": [], "total": 0}
        clean_level = _clean_dedup_level(level)
        if clean_level == "family":
            return self._fetch_dedup_family_products(
                project_name=project_name,
                category_key=category_key,
                query_text=query_text,
                limit=limit,
            )
        where_sql, params = self._dedup_products_where(
            project_name=project_name,
            category_key=category_key,
            level=clean_level,
            query_text=query_text,
        )
        safe_limit = max(1, min(int(limit), 5000))
        count = self._fetch_one(
            f"SELECT COUNT(*) AS total FROM mpstats_products_dedup {where_sql}",
            params,
            read_only=True,
            temp_directory=self._duckdb_temp_directory(),
        )
        rows = self._fetch_records(
            f"""
            SELECT {", ".join(quote_duckdb_name(column) for column in DEDUP_PRODUCTS_COLUMNS)}
            FROM mpstats_products_dedup
            {where_sql}
            ORDER BY
                category_name NULLS LAST,
                ml_family_id,
                ml_pack_id,
                sort_order,
                revenue DESC NULLS LAST,
                sku NULLS LAST
            LIMIT ?
            """,
            [*params, safe_limit],
            read_only=True,
            temp_directory=self._duckdb_temp_directory(),
        )
        return {
            "columns": list(DEDUP_PRODUCTS_COLUMNS),
            "rows": rows,
            "total": int(count["total"]) if count else 0,
        }

    def _fetch_dedup_family_products(
        self,
        *,
        project_name: str,
        category_key: str | None,
        query_text: str | None,
        limit: int,
    ) -> dict[str, Any]:
        where_sql, params = self._dedup_products_where(
            project_name=project_name,
            category_key=category_key,
            level="expanded",
            query_text=query_text,
        )
        safe_limit = max(1, min(int(limit), 5000))
        base_query = f"""
            WITH matched_rows AS (
                SELECT *
                FROM mpstats_products_dedup
                {where_sql}
            ),
            matched_packs AS (
                SELECT DISTINCT run_id, ml_family_id, ml_pack_id
                FROM matched_rows
            ),
            pack_rows AS (
                SELECT c.*
                FROM mpstats_products_dedup AS c
                JOIN matched_packs AS m
                  ON m.run_id = c.run_id
                 AND m.ml_family_id = c.ml_family_id
                 AND m.ml_pack_id = c.ml_pack_id
                WHERE c.row_level = 'canonical'
            ),
            ranked_packs AS (
                SELECT
                    p.*,
                    ROW_NUMBER() OVER (
                        PARTITION BY p.run_id, p.ml_family_id
                        ORDER BY p.revenue DESC NULLS LAST, p.source_row_count DESC NULLS LAST, p.normalized_sku NULLS LAST
                    ) AS family_rank
                FROM pack_rows AS p
            ),
            family_rows AS (
                SELECT
                    run_id,
                    MIN(project_name) AS project_name,
                    MIN(category_key) AS category_key,
                    MIN(category_name) AS category_name,
                    ml_family_id,
                    NULL AS ml_pack_id,
                    'family' AS row_level,
                    0 AS sort_order,
                    MAX(CASE WHEN family_rank = 1 THEN canonical_node_id END) AS node_id,
                    MAX(CASE WHEN family_rank = 1 THEN canonical_node_id END) AS canonical_node_id,
                    MAX(CASE WHEN family_rank = 1 THEN canonical_sku END) AS canonical_sku,
                    MAX(CASE WHEN family_rank = 1 THEN normalized_sku END) AS normalized_sku,
                    NULL AS marketplace_code,
                    NULL AS marketplace,
                    NULL AS article,
                    MAX(CASE WHEN family_rank = 1 THEN sku END) AS sku,
                    CASE
                        WHEN COUNT(DISTINCT NULLIF(brand, '')) = 1 THEN MIN(NULLIF(brand, ''))
                        ELSE NULL
                    END AS brand,
                    CASE
                        WHEN COUNT(DISTINCT NULLIF(subcategory, '')) = 1 THEN MIN(NULLIF(subcategory, ''))
                        ELSE NULL
                    END AS subcategory,
                    CASE
                        WHEN COUNT(DISTINCT unit_amount) = 1 THEN MIN(unit_amount)
                        ELSE NULL
                    END AS unit_amount,
                    NULL AS total_amount,
                    NULL AS multipack_count,
                    COALESCE(SUM(sales_volume), 0) AS sales_volume,
                    COALESCE(SUM(revenue), 0) AS revenue,
                    COALESCE(SUM(source_row_count), 0) AS source_row_count,
                    COUNT(*) AS component_size,
                    CASE WHEN COUNT(*) > 1 THEN 'family_group' ELSE 'family_singleton' END AS ml_dedup_status,
                    MAX(confidence_score) AS confidence_score
                FROM ranked_packs
                GROUP BY run_id, ml_family_id
            ),
            visible_rows AS (
                SELECT {", ".join(quote_duckdb_name(column) for column in DEDUP_PRODUCTS_COLUMNS)}
                FROM family_rows
                UNION ALL
                SELECT
                    run_id,
                    project_name,
                    category_key,
                    category_name,
                    ml_family_id,
                    ml_pack_id,
                    'pack' AS row_level,
                    1 AS sort_order,
                    node_id,
                    canonical_node_id,
                    canonical_sku,
                    normalized_sku,
                    marketplace_code,
                    marketplace,
                    article,
                    sku,
                    brand,
                    subcategory,
                    unit_amount,
                    total_amount,
                    multipack_count,
                    sales_volume,
                    revenue,
                    source_row_count,
                    component_size,
                    ml_dedup_status,
                    confidence_score
                FROM ranked_packs
            )
            SELECT {", ".join(quote_duckdb_name(column) for column in DEDUP_PRODUCTS_COLUMNS)}
            FROM visible_rows
        """
        count = self._fetch_one(
            f"SELECT COUNT(*) AS total FROM ({base_query}) AS visible_count",
            params,
            read_only=True,
            temp_directory=self._duckdb_temp_directory(),
        )
        rows = self._fetch_records(
            f"""
            {base_query}
            ORDER BY
                category_name NULLS LAST,
                ml_family_id,
                sort_order,
                unit_amount NULLS LAST,
                multipack_count NULLS LAST,
                total_amount NULLS LAST,
                revenue DESC NULLS LAST,
                normalized_sku NULLS LAST
            LIMIT ?
            """,
            [*params, safe_limit],
            read_only=True,
            temp_directory=self._duckdb_temp_directory(),
        )
        for row in rows:
            if row.get("row_level") == "family":
                title = _dedup_family_title(row.get("normalized_sku") or row.get("canonical_sku") or row.get("sku"))
                row["normalized_sku"] = title
                row["canonical_sku"] = title
                row["sku"] = title
        return {
            "columns": list(DEDUP_PRODUCTS_COLUMNS),
            "rows": rows,
            "total": int(count["total"]) if count else 0,
        }

    def export_dedup_products_csv(
        self,
        *,
        project_name: str,
        target: str | Path,
        category_key: str | None = None,
        level: str = "expanded",
    ) -> ExportResult:
        if not self.table_exists("mpstats_products_dedup"):
            raise ValueError("Таблица mpstats_products_dedup ещё не создана. Сначала запусти ML-дедуп.")
        where_sql, params = self._dedup_products_where(
            project_name=project_name,
            category_key=category_key,
            level=level,
            query_text=None,
        )
        if _clean_dedup_level(level) == "canonical":
            columns = [
                "project_name",
                "category_key",
                "category_name",
                "ml_family_id",
                "ml_pack_id",
                "normalized_sku",
                "canonical_sku",
                "brand",
                "subcategory",
                "unit_amount",
                "total_amount",
                "multipack_count",
                "sales_volume",
                "revenue",
                "source_row_count",
                "component_size",
                "run_id",
            ]
        else:
            columns = [
                "project_name",
                "category_key",
                "category_name",
                "row_level",
                "ml_family_id",
                "ml_pack_id",
                "normalized_sku",
                "canonical_sku",
                "marketplace",
                "marketplace_code",
                "article",
                "sku",
                "brand",
                "subcategory",
                "unit_amount",
                "total_amount",
                "multipack_count",
                "sales_volume",
                "revenue",
                "source_row_count",
                "component_size",
                "ml_dedup_status",
                "confidence_score",
                "run_id",
            ]
        query = f"""
            SELECT {", ".join(quote_duckdb_name(column) for column in columns)}
            FROM mpstats_products_dedup
            {where_sql}
            ORDER BY
                category_name NULLS LAST,
                ml_family_id,
                ml_pack_id,
                sort_order,
                revenue DESC NULLS LAST,
                sku NULLS LAST
        """
        return self.export_flat_query(query, Path(target), "csv", params=params, delimiter=";", header=True)

    def _dedup_products_where(
        self,
        *,
        project_name: str,
        category_key: str | None,
        level: str,
        query_text: str | None,
    ) -> tuple[str, list[Any]]:
        where = ["project_name = ?"]
        params: list[Any] = [project_name]
        if category_key:
            clean_category_key = str(category_key).strip()
            source_category_keys = self.resolve_dedup_source_category_keys(
                project_name=project_name,
                category_key=category_key,
            )
            category_keys = list(dict.fromkeys([clean_category_key, *source_category_keys]))
            for category in self.list_dedup_eligible_categories(project_name=project_name):
                group_key = str(category.get("category_key") or "").strip()
                source_keys = [str(key).strip() for key in category.get("source_category_keys") or [] if str(key).strip()]
                if clean_category_key in source_keys and group_key:
                    category_keys = list(dict.fromkeys([group_key, clean_category_key]))
                    break
            if len(category_keys) == 1:
                where.append("category_key = ?")
                params.append(category_keys[0])
            elif category_keys:
                where.append(f"category_key IN ({', '.join('?' for _ in category_keys)})")
                params.extend(category_keys)
        if _clean_dedup_level(level) == "canonical":
            where.append("row_level = 'canonical'")
        if query_text and query_text.strip():
            needle = f"%{query_text.strip().casefold()}%"
            searchable = ["normalized_sku", "canonical_sku", "sku", "brand", "article", "category_name", "marketplace"]
            where.append(
                "("
                + " OR ".join(f"lower(CAST({quote_duckdb_name(column)} AS VARCHAR)) LIKE ?" for column in searchable)
                + ")"
            )
            params.extend([needle] * len(searchable))
        return "WHERE " + " AND ".join(where), params

    def list_project_database_summaries(self, *, table_name: str) -> list[dict[str, Any]]:
        names: set[str] = set()
        summaries: dict[str, dict[str, Any]] = {}

        def ensure(project_name: str) -> dict[str, Any]:
            if project_name not in summaries:
                summaries[project_name] = {
                    "project_name": project_name,
                    "pipeline_runs_count": 0,
                    "app_runs_count": 0,
                    "tasks_count": 0,
                    "cube_slices_count": 0,
                    "cube_rows_count": 0,
                    "product_rows_count": 0,
                    "schedules_count": 0,
                    "first_period": None,
                    "latest_period": None,
                    "latest_activity": None,
                }
            return summaries[project_name]

        current_project = self.get_setting("project_name")
        if current_project:
            names.add(current_project)

        for table, column in (
            ("pipeline_runs", "project_name"),
            ("download_tasks", "project_name"),
            ("cube_registry", "project_name"),
            ("app_runs", "project_name"),
            ("app_schedules", "project_name"),
            ("pipeline_loads", "project_name"),
        ):
            for row in self._fetch_records(
                f"""
                SELECT DISTINCT {column} AS project_name
                FROM {table}
                WHERE {column} IS NOT NULL AND TRIM(CAST({column} AS VARCHAR)) <> ''
                """
            ):
                names.add(str(row["project_name"]))

        if self.table_exists(table_name):
            columns = self.table_columns(table_name)
            if "__project_name" in columns:
                quoted_table = quote_identifier(table_name)
                for row in self._fetch_records(
                    f"""
                    SELECT DISTINCT {quote_duckdb_name('__project_name')} AS project_name
                    FROM {quoted_table}
                    WHERE {quote_duckdb_name('__project_name')} IS NOT NULL
                      AND TRIM(CAST({quote_duckdb_name('__project_name')} AS VARCHAR)) <> ''
                    """
                ):
                    names.add(str(row["project_name"]))

        for name in names:
            ensure(name)

        for row in self._fetch_records(
            """
            SELECT project_name, COUNT(*) AS pipeline_runs_count, MAX(updated_at) AS latest_activity
            FROM pipeline_runs
            GROUP BY project_name
            """
        ):
            item = ensure(str(row["project_name"]))
            item["pipeline_runs_count"] = int(row["pipeline_runs_count"] or 0)
            item["latest_activity"] = row.get("latest_activity")

        for row in self._fetch_records(
            """
            SELECT project_name, COUNT(*) AS app_runs_count, MAX(COALESCE(finished_at, started_at, created_at)) AS latest_activity
            FROM app_runs
            GROUP BY project_name
            """
        ):
            item = ensure(str(row["project_name"]))
            item["app_runs_count"] = int(row["app_runs_count"] or 0)
            item["latest_activity"] = _max_iso(item.get("latest_activity"), row.get("latest_activity"))

        for row in self._fetch_records(
            """
            SELECT project_name, COUNT(*) AS tasks_count, MAX(updated_at) AS latest_activity
            FROM download_tasks
            GROUP BY project_name
            """
        ):
            item = ensure(str(row["project_name"]))
            item["tasks_count"] = int(row["tasks_count"] or 0)
            item["latest_activity"] = _max_iso(item.get("latest_activity"), row.get("latest_activity"))

        for row in self._fetch_records(
            """
            SELECT
                project_name,
                COUNT(*) AS cube_slices_count,
                SUM(rows_count) AS cube_rows_count,
                MIN(year * 12 + month) AS first_period,
                MAX(year * 12 + month) AS latest_period,
                MAX(saved_to_db_at) AS latest_activity
            FROM cube_registry
            GROUP BY project_name
            """
        ):
            item = ensure(str(row["project_name"]))
            item["cube_slices_count"] = int(row["cube_slices_count"] or 0)
            item["cube_rows_count"] = int(row["cube_rows_count"] or 0)
            item["first_period"] = _period_index_to_label(int(row["first_period"])) if row.get("first_period") else None
            item["latest_period"] = _period_index_to_label(int(row["latest_period"])) if row.get("latest_period") else None
            item["latest_activity"] = _max_iso(item.get("latest_activity"), row.get("latest_activity"))

        for row in self._fetch_records(
            """
            SELECT project_name, COUNT(*) AS schedules_count, MAX(updated_at) AS latest_activity
            FROM app_schedules
            GROUP BY project_name
            """
        ):
            item = ensure(str(row["project_name"]))
            item["schedules_count"] = int(row["schedules_count"] or 0)
            item["latest_activity"] = _max_iso(item.get("latest_activity"), row.get("latest_activity"))

        if self.table_exists(table_name):
            columns = self.table_columns(table_name)
            if "__project_name" in columns:
                quoted_table = quote_identifier(table_name)
                for row in self._fetch_records(
                    f"""
                    SELECT {quote_duckdb_name('__project_name')} AS project_name, COUNT(*) AS product_rows_count
                    FROM {quoted_table}
                    GROUP BY {quote_duckdb_name('__project_name')}
                    """
                ):
                    item = ensure(str(row["project_name"]))
                    item["product_rows_count"] = int(row["product_rows_count"] or 0)

        return sorted(summaries.values(), key=lambda item: str(item["project_name"]).casefold())

    def delete_project_records(self, *, project_name: str, table_name: str) -> dict[str, int]:
        counts: dict[str, int] = {
            "pipeline_runs": 0,
            "download_tasks": 0,
            "cube_registry": 0,
            "app_runs": 0,
            "app_run_steps": 0,
            "app_run_events": 0,
            "app_schedules": 0,
            "pipeline_loads": 0,
            "product_rows": 0,
        }

        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            counts["app_run_steps"] = int(
                con.execute(
                    """
                    SELECT COUNT(*)
                    FROM app_run_steps
                    WHERE run_id IN (SELECT run_id FROM app_runs WHERE project_name = ?)
                    """,
                    [project_name],
                ).fetchone()[0]
            )
            counts["app_run_events"] = int(
                con.execute(
                    """
                    SELECT COUNT(*)
                    FROM app_run_events
                    WHERE run_id IN (SELECT run_id FROM app_runs WHERE project_name = ?)
                    """,
                    [project_name],
                ).fetchone()[0]
            )
            for table in (
                "pipeline_runs",
                "download_tasks",
                "cube_registry",
                "app_runs",
                "app_schedules",
                "pipeline_loads",
            ):
                counts[table] = int(con.execute(f"SELECT COUNT(*) FROM {table} WHERE project_name = ?", [project_name]).fetchone()[0])

            quoted_table = quote_identifier(table_name)
            products_exists = bool(
                con.execute(
                    """
                    SELECT COUNT(*)
                    FROM information_schema.tables
                    WHERE table_schema = 'main' AND table_name = ?
                    """,
                    [table_name],
                ).fetchone()[0]
            )
            if products_exists:
                product_columns = {
                    str(row[0])
                    for row in con.execute(
                        """
                        SELECT column_name
                        FROM information_schema.columns
                        WHERE table_schema = 'main' AND table_name = ?
                        """,
                        [table_name],
                    ).fetchall()
                }
                if "__project_name" in product_columns:
                    counts["product_rows"] = int(
                        con.execute(
                            f"SELECT COUNT(*) FROM {quoted_table} WHERE {quote_duckdb_name('__project_name')} = ?",
                            [project_name],
                        ).fetchone()[0]
                    )

            con.execute(
                """
                DELETE FROM app_run_steps
                WHERE run_id IN (SELECT run_id FROM app_runs WHERE project_name = ?)
                """,
                [project_name],
            )
            con.execute(
                """
                DELETE FROM app_run_events
                WHERE run_id IN (SELECT run_id FROM app_runs WHERE project_name = ?)
                """,
                [project_name],
            )
            for table in (
                "download_tasks",
                "pipeline_runs",
                "cube_registry",
                "app_schedules",
                "pipeline_loads",
                "app_runs",
            ):
                con.execute(f"DELETE FROM {table} WHERE project_name = ?", [project_name])

            if products_exists and counts["product_rows"]:
                con.execute(
                    f"DELETE FROM {quoted_table} WHERE {quote_duckdb_name('__project_name')} = ?",
                    [project_name],
                )

        return counts

    def delete_pipeline_run(self, run_id: str) -> dict[str, int]:
        counts = {"pipeline_runs": 0, "download_tasks": 0}
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            counts["pipeline_runs"] = int(con.execute("SELECT COUNT(*) FROM pipeline_runs WHERE id = ?", [run_id]).fetchone()[0])
            counts["download_tasks"] = int(con.execute("SELECT COUNT(*) FROM download_tasks WHERE run_id = ?", [run_id]).fetchone()[0])
            con.execute("DELETE FROM download_tasks WHERE run_id = ?", [run_id])
            con.execute("DELETE FROM pipeline_runs WHERE id = ?", [run_id])
        return counts

    def upsert_category(self, category: dict[str, Any]) -> None:
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute(
                """
                INSERT INTO app_category_catalog (
                    category_id, category_name, marketplace, mp_code, source_type, path,
                    filter_json, fbs, period_from, period_to, source_file, is_active, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, now())
                ON CONFLICT (category_id) DO UPDATE SET
                    category_name = EXCLUDED.category_name,
                    marketplace = EXCLUDED.marketplace,
                    mp_code = EXCLUDED.mp_code,
                    source_type = EXCLUDED.source_type,
                    path = EXCLUDED.path,
                    filter_json = EXCLUDED.filter_json,
                    fbs = EXCLUDED.fbs,
                    period_from = EXCLUDED.period_from,
                    period_to = EXCLUDED.period_to,
                    source_file = EXCLUDED.source_file,
                    is_active = EXCLUDED.is_active,
                    updated_at = now()
                """,
                [
                    category["category_id"],
                    category["category_name"],
                    category["marketplace"],
                    category["mp_code"],
                    category.get("source_type") or "category",
                    category["path"],
                    category.get("filter_json"),
                    category.get("fbs"),
                    category.get("period_from"),
                    category.get("period_to"),
                    category.get("source_file"),
                    bool(category.get("is_active", True)),
                ],
            )

    def replace_categories(self, categories: list[dict[str, Any]], *, source_file: str) -> None:
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute("UPDATE app_category_catalog SET is_active = false, updated_at = now()")
        for category in categories:
            self.upsert_category(category)

    def list_categories(self, *, active_only: bool = True) -> list[dict[str, Any]]:
        query = """
            SELECT *
            FROM app_category_catalog
        """
        if active_only:
            query += " WHERE is_active = true"
        query += " ORDER BY category_name, marketplace, period_from NULLS FIRST, period_to NULLS LAST, path"
        return self._fetch_records(query)

    def get_categories_by_ids(self, category_ids: list[str]) -> list[dict[str, Any]]:
        if not category_ids:
            return []
        placeholders = ", ".join("?" for _ in category_ids)
        return self._fetch_records(
            f"""
            SELECT *
            FROM app_category_catalog
            WHERE category_id IN ({placeholders})
            ORDER BY category_name, marketplace, period_from NULLS FIRST, period_to NULLS LAST, path
            """,
            list(category_ids),
        )

    def create_pipeline_run(
        self,
        *,
        run_id: str,
        project_name: str,
        run_type: str,
        period_from: str,
        period_to: str,
        selected_category_ids: list[str],
        settings: dict[str, Any],
    ) -> dict[str, Any]:
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute(
                """
                INSERT INTO pipeline_runs (
                    id, project_name, run_type, period_from, period_to, status,
                    selected_category_ids_json, settings_json, updated_at
                )
                VALUES (?, ?, ?, ?, ?, 'planned', ?, ?, now())
                """,
                [
                    run_id,
                    project_name,
                    run_type,
                    period_from,
                    period_to,
                    json.dumps(selected_category_ids, ensure_ascii=False),
                    json.dumps(settings, ensure_ascii=False),
                ],
            )
        return self.get_pipeline_run(run_id) or {}

    def get_pipeline_run(self, run_id: str) -> dict[str, Any] | None:
        return self._fetch_one("SELECT * FROM pipeline_runs WHERE id = ?", [run_id])

    def list_pipeline_runs(self, *, project_name: str | None = None, limit: int = 20) -> list[dict[str, Any]]:
        query = "SELECT * FROM pipeline_runs"
        params: list[Any] = []
        if project_name:
            query += " WHERE project_name = ?"
            params.append(project_name)
        query += " ORDER BY created_at DESC LIMIT ?"
        params.append(int(limit))
        return self._fetch_records(query, params)

    def update_pipeline_run(self, run_id: str, values: dict[str, Any]) -> dict[str, Any] | None:
        allowed = {
            "status",
            "total_tasks",
            "completed_tasks",
            "failed_tasks",
            "current_step",
            "pause_requested",
            "stop_requested",
            "started_at",
            "finished_at",
        }
        assignments = [key for key in values if key in allowed]
        if not assignments:
            return self.get_pipeline_run(run_id)
        sql = ", ".join(f"{key} = ?" for key in assignments)
        params = [values[key] for key in assignments]
        params.append(run_id)
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute(f"UPDATE pipeline_runs SET {sql}, updated_at = now() WHERE id = ?", params)
        return self.get_pipeline_run(run_id)

    def request_pipeline_pause(self, run_id: str) -> dict[str, Any] | None:
        return self.update_pipeline_run(run_id, {"pause_requested": True, "status": "pausing"})

    def request_pipeline_stop(self, run_id: str) -> dict[str, Any] | None:
        return self.update_pipeline_run(run_id, {"stop_requested": True, "pause_requested": False, "status": "stopping"})

    def clear_pipeline_pause(self, run_id: str) -> dict[str, Any] | None:
        return self.update_pipeline_run(run_id, {"pause_requested": False})

    def clear_pipeline_control(self, run_id: str) -> dict[str, Any] | None:
        return self.update_pipeline_run(run_id, {"pause_requested": False, "stop_requested": False})

    def is_pipeline_pause_requested(self, run_id: str) -> bool:
        row = self._fetch_one("SELECT pause_requested FROM pipeline_runs WHERE id = ?", [run_id])
        return bool(row and row.get("pause_requested"))

    def is_pipeline_stop_requested(self, run_id: str) -> bool:
        row = self._fetch_one("SELECT stop_requested FROM pipeline_runs WHERE id = ?", [run_id])
        return bool(row and row.get("stop_requested"))

    def refresh_pipeline_run_counts(self, run_id: str) -> dict[str, Any] | None:
        row = self._fetch_one(
            """
            SELECT
                COUNT(*) AS total_tasks,
                SUM(CASE WHEN status IN ('saved_to_db', 'skipped', 'no_data') THEN 1 ELSE 0 END) AS completed_tasks,
                SUM(CASE WHEN status = 'failed' THEN 1 ELSE 0 END) AS failed_tasks
            FROM download_tasks
            WHERE run_id = ?
            """,
            [run_id],
        )
        return self.update_pipeline_run(
            run_id,
            {
                "total_tasks": int(row["total_tasks"] or 0) if row else 0,
                "completed_tasks": int(row["completed_tasks"] or 0) if row else 0,
                "failed_tasks": int(row["failed_tasks"] or 0) if row else 0,
            },
        )

    def upsert_download_task(self, task: dict[str, Any]) -> dict[str, Any]:
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute(
                """
                INSERT INTO download_tasks (
                    id, run_id, project_name, marketplace, marketplace_code,
                    source_type, category_name, category_path, category_id, category_key,
                    year, month, status, download_status, process_status,
                    classify_status, save_status, raw_file_path, processed_file_path,
                    classified_file_path, rows_count, error_message, task_hash, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, now())
                ON CONFLICT (project_name, marketplace_code, category_key, year, month) DO UPDATE SET
                    run_id = EXCLUDED.run_id,
                    marketplace = EXCLUDED.marketplace,
                    source_type = EXCLUDED.source_type,
                    category_name = EXCLUDED.category_name,
                    category_path = EXCLUDED.category_path,
                    category_id = EXCLUDED.category_id,
                    status = EXCLUDED.status,
                    download_status = EXCLUDED.download_status,
                    process_status = EXCLUDED.process_status,
                    classify_status = EXCLUDED.classify_status,
                    save_status = EXCLUDED.save_status,
                    raw_file_path = EXCLUDED.raw_file_path,
                    processed_file_path = EXCLUDED.processed_file_path,
                    classified_file_path = EXCLUDED.classified_file_path,
                    rows_count = EXCLUDED.rows_count,
                    error_message = EXCLUDED.error_message,
                    task_hash = EXCLUDED.task_hash,
                    updated_at = now()
                """,
                [
                    task["id"],
                    task["run_id"],
                    task["project_name"],
                    task["marketplace"],
                    task["marketplace_code"],
                    task.get("source_type") or "category",
                    task["category_name"],
                    task["category_path"],
                    task["category_id"],
                    task["category_key"],
                    int(task["year"]),
                    int(task["month"]),
                    task["status"],
                    task["download_status"],
                    task["process_status"],
                    task["classify_status"],
                    task["save_status"],
                    task.get("raw_file_path"),
                    task.get("processed_file_path"),
                    task.get("classified_file_path"),
                    int(task.get("rows_count") or 0),
                    task.get("error_message"),
                    task["task_hash"],
                ],
            )
        return self.get_download_task(str(task["id"])) or {}

    def get_download_task(self, task_id: str) -> dict[str, Any] | None:
        return self._fetch_one("SELECT * FROM download_tasks WHERE id = ?", [task_id])

    def list_download_tasks(self, *, run_id: str, task_filter: str = "all", limit: int | None = None) -> list[dict[str, Any]]:
        query = "SELECT * FROM download_tasks WHERE run_id = ?"
        params: list[Any] = [run_id]
        if task_filter == "errors":
            query += " AND status = 'failed'"
        elif task_filter == "not_downloaded":
            query += " AND download_status NOT IN ('downloaded', 'skipped')"
        elif task_filter == "not_processed":
            query += " AND process_status <> 'processed'"
        elif task_filter == "not_saved":
            query += " AND save_status <> 'saved_to_db'"
        elif task_filter == "ready":
            query += " AND status IN ('processed', 'classified', 'saved_to_db')"
        query += " ORDER BY year, month, category_name, marketplace"
        if limit is not None:
            query += " LIMIT ?"
            params.append(int(limit))
        return self._fetch_records(query, params)

    def summarize_download_tasks(self, *, run_id: str) -> dict[str, int]:
        row = self._fetch_one(
            """
            SELECT
                COUNT(*) AS total_tasks,
                COUNT(DISTINCT category_id) AS category_count,
                COUNT(DISTINCT CAST(year AS VARCHAR) || '-' || CAST(month AS VARCHAR)) AS month_count
            FROM download_tasks
            WHERE run_id = ?
            """,
            [run_id],
        )
        return {
            "total_tasks": int(row["total_tasks"] or 0) if row else 0,
            "category_count": int(row["category_count"] or 0) if row else 0,
            "month_count": int(row["month_count"] or 0) if row else 0,
        }

    def list_project_download_tasks(self, *, project_name: str, limit: int = 500) -> list[dict[str, Any]]:
        return self._fetch_records(
            """
            SELECT *
            FROM download_tasks
            WHERE project_name = ?
            ORDER BY updated_at DESC
            LIMIT ?
            """,
            [project_name, int(limit)],
        )

    def update_download_task(self, task_id: str, values: dict[str, Any]) -> dict[str, Any] | None:
        allowed = {
            "status",
            "download_status",
            "process_status",
            "classify_status",
            "save_status",
            "raw_file_path",
            "processed_file_path",
            "classified_file_path",
            "rows_count",
            "error_message",
        }
        assignments = [key for key in values if key in allowed]
        if not assignments:
            return self.get_download_task(task_id)
        sql = ", ".join(f"{key} = ?" for key in assignments)
        params = [values[key] for key in assignments]
        params.append(task_id)
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute(f"UPDATE download_tasks SET {sql}, updated_at = now() WHERE id = ?", params)
        task = self.get_download_task(task_id)
        if task:
            self.refresh_pipeline_run_counts(str(task["run_id"]))
        return task

    def reset_failed_tasks(self, run_id: str) -> int:
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            count = con.execute(
                "SELECT COUNT(*) FROM download_tasks WHERE run_id = ? AND status = 'failed'",
                [run_id],
            ).fetchone()[0]
            con.execute(
                """
                UPDATE download_tasks
                SET status = 'pending', error_message = NULL, updated_at = now()
                WHERE run_id = ? AND status = 'failed'
                """,
                [run_id],
            )
        self.refresh_pipeline_run_counts(run_id)
        return int(count)

    def reset_task_for_retry(self, task_id: str) -> dict[str, Any] | None:
        return self.update_download_task(task_id, {"status": "pending", "error_message": None})

    def get_cube_entry(
        self,
        *,
        project_name: str,
        year: int,
        month: int,
        marketplace_code: str,
        category_key: str,
    ) -> dict[str, Any] | None:
        return self._fetch_one(
            """
            SELECT *
            FROM cube_registry
            WHERE project_name = ? AND year = ? AND month = ? AND marketplace_code = ? AND category_key = ?
            """,
            [project_name, int(year), int(month), marketplace_code, category_key],
        )

    def list_cube_entries_for_tasks(self, tasks: list[dict[str, Any]]) -> dict[tuple[str, int, int, str, str], dict[str, Any]]:
        keys = {
            (
                str(task["project_name"]),
                int(task["year"]),
                int(task["month"]),
                str(task["marketplace_code"]),
                str(task["category_key"]),
            )
            for task in tasks
        }
        if not keys:
            return {}

        projects = sorted({key[0] for key in keys})
        category_keys = sorted({key[4] for key in keys})
        project_placeholders = ", ".join("?" for _ in projects)
        category_placeholders = ", ".join("?" for _ in category_keys)
        rows = self._fetch_records(
            f"""
            SELECT *
            FROM cube_registry
            WHERE project_name IN ({project_placeholders})
              AND category_key IN ({category_placeholders})
            """,
            [*projects, *category_keys],
        )
        out: dict[tuple[str, int, int, str, str], dict[str, Any]] = {}
        for row in rows:
            key = (
                str(row["project_name"]),
                int(row["year"]),
                int(row["month"]),
                str(row["marketplace_code"]),
                str(row["category_key"]),
            )
            if key in keys:
                out[key] = row
        return out

    def get_cube_entry_by_natural_key(
        self,
        *,
        project_name: str,
        year: int,
        month: int,
        marketplace_code: str,
        category_name: str,
    ) -> dict[str, Any] | None:
        return self._fetch_one(
            """
            SELECT *
            FROM cube_registry
            WHERE project_name = ?
              AND year = ?
              AND month = ?
              AND marketplace_code = ?
              AND lower(trim(CAST(category_name AS VARCHAR))) = lower(trim(CAST(? AS VARCHAR)))
            ORDER BY saved_to_db_at DESC
            LIMIT 1
            """,
            [project_name, int(year), int(month), marketplace_code, category_name],
        )

    def get_cube_entry_by_id(self, entry_id: str) -> dict[str, Any] | None:
        return self._fetch_one("SELECT * FROM cube_registry WHERE id = ?", [entry_id])

    def upsert_cube_entry(self, entry: dict[str, Any]) -> dict[str, Any]:
        entry_id = str(entry.get("id") or uuid4().hex)
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute(
                """
                INSERT INTO cube_registry (
                    id, project_name, year, month, marketplace, marketplace_code,
                    source_type, category_key, category_name, rows_count, saved_to_db_at,
                    exported_at,
                    source_processed_file_path, file_hash, days_loaded,
                    days_in_month, data_actual_until, data_mode, is_heavy,
                    heavy_reason
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, now(), COALESCE(?, now()), ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT (project_name, year, month, marketplace_code, category_key) DO UPDATE SET
                    marketplace = EXCLUDED.marketplace,
                    source_type = EXCLUDED.source_type,
                    category_name = EXCLUDED.category_name,
                    rows_count = EXCLUDED.rows_count,
                    saved_to_db_at = now(),
                    exported_at = EXCLUDED.exported_at,
                    source_processed_file_path = EXCLUDED.source_processed_file_path,
                    file_hash = EXCLUDED.file_hash,
                    days_loaded = EXCLUDED.days_loaded,
                    days_in_month = EXCLUDED.days_in_month,
                    data_actual_until = EXCLUDED.data_actual_until,
                    data_mode = EXCLUDED.data_mode,
                    is_heavy = EXCLUDED.is_heavy,
                    heavy_reason = EXCLUDED.heavy_reason
                """,
                [
                    entry_id,
                    entry["project_name"],
                    int(entry["year"]),
                    int(entry["month"]),
                    entry["marketplace"],
                    entry["marketplace_code"],
                    entry.get("source_type") or "category",
                    entry["category_key"],
                    entry["category_name"],
                    int(entry.get("rows_count") or 0),
                    entry.get("exported_at"),
                    entry.get("source_processed_file_path"),
                    entry.get("file_hash"),
                    entry.get("days_loaded"),
                    entry.get("days_in_month"),
                    entry.get("data_actual_until"),
                    entry.get("data_mode") or "standard",
                    bool(entry.get("is_heavy", False)),
                    entry.get("heavy_reason"),
                ],
            )
        return self.get_cube_entry(
            project_name=str(entry["project_name"]),
            year=int(entry["year"]),
            month=int(entry["month"]),
            marketplace_code=str(entry["marketplace_code"]),
            category_key=str(entry["category_key"]),
        ) or {}

    def list_cube_registry(self, *, project_name: str, limit: int = 500) -> list[dict[str, Any]]:
        self.refresh_large_category_flags(project_name=project_name)
        return self._fetch_records(
            """
            SELECT *
            FROM cube_registry
            WHERE project_name = ?
            ORDER BY year DESC, month DESC, category_name, marketplace
            LIMIT ?
            """,
            [project_name, int(limit)],
        )

    def count_cube_registry(self, *, project_name: str) -> int:
        row = self._fetch_one(
            "SELECT COUNT(*) AS total FROM cube_registry WHERE project_name = ?",
            [project_name],
        )
        return int(row["total"] if row else 0)

    def list_cube_registry_by_source_file(self, *, project_name: str, source_file_path: str) -> list[dict[str, Any]]:
        return self._fetch_records(
            """
            SELECT *
            FROM cube_registry
            WHERE project_name = ? AND source_processed_file_path = ?
            ORDER BY year DESC, month DESC, category_name, marketplace
            """,
            [project_name, source_file_path],
        )

    def delete_cube_entry(self, *, entry_id: str, table_name: str) -> dict[str, Any]:
        result = self.delete_cube_entries(entry_ids=[entry_id], table_name=table_name)
        entry = result["entries"][0] if result["entries"] else None
        return {"entry_id": entry_id, "deleted": result["deleted"], "entry": entry}

    def delete_cube_entries(self, *, entry_ids: list[str], table_name: str) -> dict[str, Any]:
        unique_ids = list(dict.fromkeys(str(entry_id).strip() for entry_id in entry_ids if str(entry_id).strip()))
        counts: dict[str, int] = {"cube_registry": 0, "product_rows": 0, "download_tasks": 0}
        run_ids: set[str] = set()
        deleted_entries: list[dict[str, Any]] = []
        affected_categories: dict[str, set[str]] = {}
        if not unique_ids:
            return {"entry_ids": [], "deleted": counts, "entries": []}

        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            placeholders = ", ".join("?" for _ in unique_ids)
            result = con.execute(
                f"""
                SELECT *
                FROM cube_registry
                WHERE id IN ({placeholders})
                """,
                unique_ids,
            )
            columns = [col[0] for col in result.description]
            deleted_entries = [clean_record(dict(zip(columns, row))) for row in result.fetchall()]
            if not deleted_entries:
                return {"entry_ids": unique_ids, "deleted": counts, "entries": []}

            con.execute("DROP TABLE IF EXISTS _cube_delete_keys")
            con.execute(
                f"""
                CREATE TEMPORARY TABLE _cube_delete_keys AS
                SELECT id, project_name, year, month, marketplace_code, category_key
                FROM cube_registry
                WHERE id IN ({placeholders})
                """,
                unique_ids,
            )

            if table_exists(con, table_name):
                product_columns = set(_table_column_types(con, table_name))
                required = {"__project_name", "__year", "__month", "__marketplace_code", "__category_key"}
                if required.issubset(product_columns):
                    quoted_table = quote_identifier(table_name)
                    product_join_sql = f"""
                        p.{quote_duckdb_name('__project_name')} = k.project_name
                        AND CAST(p.{quote_duckdb_name('__year')} AS INTEGER) = k.year
                        AND CAST(p.{quote_duckdb_name('__month')} AS INTEGER) = k.month
                        AND p.{quote_duckdb_name('__marketplace_code')} = k.marketplace_code
                        AND p.{quote_duckdb_name('__category_key')} = k.category_key
                    """
                    counts["product_rows"] = int(
                        con.execute(
                            f"""
                            SELECT COUNT(*)
                            FROM {quoted_table} AS p
                            JOIN _cube_delete_keys AS k ON {product_join_sql}
                            """
                        ).fetchone()[0]
                    )
                    if counts["product_rows"]:
                        con.execute(
                            f"""
                            DELETE FROM {quoted_table} AS p
                            USING _cube_delete_keys AS k
                            WHERE {product_join_sql}
                            """
                        )

            run_ids = {
                str(task_row[0])
                for task_row in con.execute(
                    """
                    SELECT DISTINCT d.run_id
                    FROM download_tasks AS d
                    JOIN _cube_delete_keys AS k
                      ON d.project_name = k.project_name
                     AND d.year = k.year
                     AND d.month = k.month
                     AND d.marketplace_code = k.marketplace_code
                     AND d.category_key = k.category_key
                    """
                ).fetchall()
                if task_row[0] is not None
            }
            counts["download_tasks"] = int(
                con.execute(
                    """
                    SELECT COUNT(*)
                    FROM download_tasks AS d
                    JOIN _cube_delete_keys AS k
                      ON d.project_name = k.project_name
                     AND d.year = k.year
                     AND d.month = k.month
                     AND d.marketplace_code = k.marketplace_code
                     AND d.category_key = k.category_key
                    """
                ).fetchone()[0]
            )
            con.execute(
                """
                UPDATE download_tasks AS d
                SET
                    save_status = 'pending',
                    status = CASE
                        WHEN classify_status = 'classified' THEN 'classified'
                        WHEN process_status = 'processed' THEN 'processed'
                        WHEN download_status = 'downloaded' THEN 'downloaded'
                        ELSE 'pending'
                    END,
                    error_message = NULL,
                    updated_at = now()
                FROM _cube_delete_keys AS k
                WHERE d.project_name = k.project_name
                  AND d.year = k.year
                  AND d.month = k.month
                  AND d.marketplace_code = k.marketplace_code
                  AND d.category_key = k.category_key
                """
            )
            counts["cube_registry"] = int(
                con.execute("SELECT COUNT(*) FROM cube_registry WHERE id IN ({})".format(placeholders), unique_ids).fetchone()[0]
            )
            con.execute(f"DELETE FROM cube_registry WHERE id IN ({placeholders})", unique_ids)

            for entry in deleted_entries:
                affected_categories.setdefault(str(entry["project_name"]), set()).add(str(entry["category_key"]))

        for run_id in run_ids:
            self.refresh_pipeline_run_counts(run_id)
        for project_name, category_keys in affected_categories.items():
            self.refresh_large_category_flags(project_name=project_name, category_keys=sorted(category_keys))
        return {"entry_ids": unique_ids, "deleted": counts, "entries": deleted_entries}

    def mark_project_file_deleted(self, *, project_name: str, file_path: str, file_kind: str) -> dict[str, int]:
        if file_kind not in {"raw", "processed", "classified"}:
            return {"download_tasks": 0}

        path_column = {
            "raw": "raw_file_path",
            "processed": "processed_file_path",
            "classified": "classified_file_path",
        }[file_kind]
        run_ids: set[str] = set()
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            run_ids = {
                str(task_row[0])
                for task_row in con.execute(
                    f"""
                    SELECT DISTINCT run_id
                    FROM download_tasks
                    WHERE project_name = ? AND {path_column} = ?
                    """,
                    [project_name, file_path],
                ).fetchall()
                if task_row[0] is not None
            }
            count = int(
                con.execute(
                    f"""
                    SELECT COUNT(*)
                    FROM download_tasks
                    WHERE project_name = ? AND {path_column} = ?
                    """,
                    [project_name, file_path],
                ).fetchone()[0]
            )
            if not count:
                return {"download_tasks": 0}

            if file_kind == "raw":
                con.execute(
                    """
                    UPDATE download_tasks
                    SET
                        download_status = 'pending',
                        status = CASE
                            WHEN save_status = 'saved_to_db' THEN status
                            WHEN classify_status = 'classified' THEN status
                            WHEN process_status = 'processed' THEN status
                            ELSE 'pending'
                        END,
                        updated_at = now()
                    WHERE project_name = ? AND raw_file_path = ?
                    """,
                    [project_name, file_path],
                )
            elif file_kind == "processed":
                con.execute(
                    """
                    UPDATE download_tasks
                    SET
                        process_status = 'pending',
                        classify_status = CASE WHEN classify_status = 'classified' THEN classify_status ELSE 'pending' END,
                        save_status = CASE WHEN save_status = 'saved_to_db' THEN save_status ELSE 'pending' END,
                        status = CASE
                            WHEN save_status = 'saved_to_db' THEN status
                            WHEN classify_status = 'classified' THEN status
                            WHEN download_status = 'downloaded' THEN 'downloaded'
                            ELSE 'pending'
                        END,
                        updated_at = now()
                    WHERE project_name = ? AND processed_file_path = ?
                    """,
                    [project_name, file_path],
                )
            else:
                con.execute(
                    """
                    UPDATE download_tasks
                    SET
                        classify_status = 'pending',
                        save_status = 'pending',
                        status = CASE
                            WHEN process_status = 'processed' THEN 'processed'
                            WHEN download_status = 'downloaded' THEN 'downloaded'
                            ELSE 'pending'
                        END,
                        updated_at = now()
                    WHERE project_name = ? AND classified_file_path = ?
                    """,
                    [project_name, file_path],
                )

        for run_id in run_ids:
            self.refresh_pipeline_run_counts(run_id)
        return {"download_tasks": count}

    def latest_cube_month(self, *, project_name: str) -> tuple[int, int] | None:
        row = self._fetch_one(
            """
            SELECT year, month
            FROM cube_registry
            WHERE project_name = ?
            ORDER BY year DESC, month DESC
            LIMIT 1
            """,
            [project_name],
        )
        if not row:
            return None
        return int(row["year"]), int(row["month"])

    def latest_successful_run_id(self, *, project_name: str | None = None) -> str | None:
        app_where = "status = 'succeeded'"
        pipeline_where = "status IN ('succeeded', 'completed_with_errors')"
        params: list[Any] = []
        if project_name:
            app_where += " AND project_name = ?"
            pipeline_where += " AND project_name = ?"
            params.extend([project_name, project_name])
        row = self._fetch_one(
            f"""
            SELECT run_id
            FROM (
                SELECT run_id, finished_at, created_at
                FROM app_runs
                WHERE {app_where}
                UNION ALL
                SELECT id AS run_id, finished_at, created_at
                FROM pipeline_runs
                WHERE {pipeline_where}
            )
            ORDER BY finished_at DESC NULLS LAST, created_at DESC
            LIMIT 1
            """,
            params,
        )
        return str(row["run_id"]) if row else None

    def mark_run_running(self, run_id: str) -> None:
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute(
                "UPDATE app_runs SET status = 'running', started_at = now() WHERE run_id = ?",
                [run_id],
            )

    def finish_run(self, run_id: str, status: str, *, error: str | None = None, manifest_path: str | None = None) -> None:
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute(
                """
                UPDATE app_runs
                SET status = ?, error = ?, manifest_path = COALESCE(?, manifest_path), finished_at = now()
                WHERE run_id = ?
                """,
                [status, error, manifest_path, run_id],
            )

    def request_cancel(self, run_id: str) -> None:
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute("UPDATE app_runs SET requested_cancel = true WHERE run_id = ?", [run_id])
        self.add_event(run_id, "warning", "Запрошена отмена прогона", None)

    def is_cancel_requested(self, run_id: str) -> bool:
        row = self._fetch_one("SELECT requested_cancel FROM app_runs WHERE run_id = ?", [run_id])
        return bool(row and row.get("requested_cancel"))

    def record_step(
        self,
        *,
        run_id: str,
        step_number: int,
        step_name: str,
        status: str,
        rows: int = 0,
        ok_count: int = 0,
        error_count: int = 0,
        skipped_count: int = 0,
        output: str | None = None,
        details: list[dict[str, Any]] | None = None,
        error: str | None = None,
        finished: bool = False,
    ) -> None:
        details_json = json.dumps(details or [], ensure_ascii=False)
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute(
                "DELETE FROM app_run_steps WHERE run_id = ? AND step_number = ?",
                [run_id, step_number],
            )
            con.execute(
                """
                INSERT INTO app_run_steps (
                    run_id, step_number, step_name, status, rows_loaded, ok_count,
                    error_count, skipped_count, output, details_json, error, finished_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CASE WHEN ? THEN now() ELSE NULL END)
                """,
                [
                    run_id,
                    step_number,
                    step_name,
                    status,
                    rows,
                    ok_count,
                    error_count,
                    skipped_count,
                    output,
                    details_json,
                    error,
                    finished,
                ],
            )

    def list_run_steps(self, run_id: str) -> list[dict[str, Any]]:
        return self._fetch_records(
            """
            SELECT *
            FROM app_run_steps
            WHERE run_id = ?
            ORDER BY step_number
            """,
            [run_id],
        )

    def add_event(
        self,
        run_id: str,
        level: str,
        message: str,
        payload: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        event_id = uuid4().hex
        payload_json = json.dumps(payload, ensure_ascii=False) if payload else None
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute(
                """
                INSERT INTO app_run_events (event_id, run_id, level, message, payload_json)
                VALUES (?, ?, ?, ?, ?)
                """,
                [event_id, run_id, level, message, payload_json],
            )
        return self._fetch_one("SELECT * FROM app_run_events WHERE event_id = ?", [event_id]) or {}

    def list_run_events(self, run_id: str, *, after: str | None = None, limit: int = 200) -> list[dict[str, Any]]:
        query = """
            SELECT *
            FROM app_run_events
            WHERE run_id = ?
        """
        params: list[Any] = [run_id]
        if after:
            query += " AND created_at > ?"
            params.append(after)
        query += " ORDER BY created_at ASC LIMIT ?"
        params.append(int(limit))
        return self._fetch_records(query, params)

    def table_columns(self, table_name: str) -> list[str]:
        quote_identifier(table_name)
        rows = self._fetch_records(
            """
            SELECT column_name
            FROM information_schema.columns
            WHERE table_schema = 'main' AND table_name = ?
            ORDER BY ordinal_position
            """,
            [table_name],
        )
        return [str(row["column_name"]) for row in rows]

    def table_exists(self, table_name: str) -> bool:
        quote_identifier(table_name)
        row = self._fetch_one(
            """
            SELECT COUNT(*) AS cnt
            FROM information_schema.tables
            WHERE table_schema = 'main' AND table_name = ?
            """,
            [table_name],
        )
        return bool(row and int(row["cnt"]) > 0)

    def import_products_file(
        self,
        *,
        run_id: str,
        csv_path: Path,
        table_name: str,
        project_name: str,
        load_name: str | None = None,
    ) -> int:
        quoted_table = quote_identifier(table_name)
        stage_table = "_mpstats_app_products_stage"
        rows_loaded = 0
        with measure_duckdb_operation(
            "import_products_file",
            {"db_path": self.settings.db_path, "file_path": csv_path, "table_name": table_name, "project_name": project_name},
        ) as metrics:
            with self._lock, connect(self.settings.db_path, temp_directory=self._duckdb_temp_directory()) as con:
                apply_migrations(con)
                with duckdb_transaction(con):
                    stage_columns = _create_products_stage(
                        con,
                        csv_path=csv_path,
                        stage_table=stage_table,
                        run_id=run_id,
                        source_file=str(csv_path),
                        project_name=project_name,
                    )
                    rows_loaded = _stage_count(con, stage_table)
                    exists = table_exists(con, table_name)
                    if not exists:
                        con.execute(f"CREATE TABLE {quoted_table} AS SELECT * FROM {quote_identifier(stage_table)}")
                    else:
                        target_types = _ensure_table_accepts_stage_columns(
                            con,
                            table_name=table_name,
                            quoted_table=quoted_table,
                            stage_table=stage_table,
                            stage_columns=stage_columns,
                        )
                        con.execute(
                            _insert_stage_sql(
                                quoted_table=quoted_table,
                                stage_table=stage_table,
                                columns=stage_columns,
                                target_types=target_types,
                                deduplicate=False,
                            )
                        )

                    con.execute(
                        """
                        INSERT INTO pipeline_loads (
                            table_name, source_file, load_name, project_name, mode, rows_loaded
                        )
                        VALUES (?, ?, ?, ?, 'append', ?)
                        """,
                        [table_name, str(csv_path), load_name or f"app_run:{run_id}", project_name, rows_loaded],
                    )
            metrics["rows_affected"] = rows_loaded
        return rows_loaded

    def import_products_file_idempotent(
        self,
        *,
        run_id: str,
        csv_path: Path,
        table_name: str,
        project_name: str,
        year: int,
        month: int,
        marketplace_code: str,
        category_key: str,
        category_name: str | None = None,
        source_type: str = "category",
        overwrite: bool = False,
        load_name: str | None = None,
    ) -> int:
        quoted_table = quote_identifier(table_name)
        inserted = 0
        stage_table = "_mpstats_app_products_stage"
        with measure_duckdb_operation(
            "import_products_file_idempotent",
            {
                "db_path": self.settings.db_path,
                "file_path": csv_path,
                "table_name": table_name,
                "project_name": project_name,
                "year": int(year),
                "month": int(month),
                "marketplace_code": marketplace_code,
                "category_key": category_key,
            },
        ) as metrics:
            with self._lock, connect(self.settings.db_path, temp_directory=self._duckdb_temp_directory()) as con:
                apply_migrations(con)
                with duckdb_transaction(con):
                    stage_columns = _create_products_stage(
                        con,
                        csv_path=csv_path,
                        stage_table=stage_table,
                        run_id=run_id,
                        source_file=str(csv_path),
                        project_name=project_name,
                        year=int(year),
                        month=int(month),
                        marketplace_code=marketplace_code,
                        source_type=source_type,
                        category_key=category_key,
                    )
                    exists = table_exists(con, table_name)
                    effective_category_name = category_name
                    delete_natural_slice = False
                    if not effective_category_name and "Категория" in stage_columns:
                        effective_category_name = _stage_single_text_value(con, stage_table=stage_table, column="Категория")
                    if effective_category_name:
                        registry_conflict = _fetch_cube_entry_by_natural_key(
                            con,
                            project_name=project_name,
                            year=int(year),
                            month=int(month),
                            marketplace_code=marketplace_code,
                            category_name=effective_category_name,
                        )
                        if registry_conflict and not overwrite:
                            raise DuplicateCubeSliceError(
                                _duplicate_cube_slice_message(
                                    project_name=project_name,
                                    year=int(year),
                                    month=int(month),
                                    marketplace_code=marketplace_code,
                                    category_name=effective_category_name,
                                )
                            )
                        if exists:
                            product_conflict_rows = _product_slice_count_by_natural_key(
                                con,
                                table_name=table_name,
                                quoted_table=quoted_table,
                                project_name=project_name,
                                year=int(year),
                                month=int(month),
                                marketplace_code=marketplace_code,
                                category_name=effective_category_name,
                            )
                            if product_conflict_rows and not overwrite:
                                raise DuplicateCubeSliceError(
                                    _duplicate_cube_slice_message(
                                        project_name=project_name,
                                        year=int(year),
                                        month=int(month),
                                        marketplace_code=marketplace_code,
                                        category_name=effective_category_name,
                                    )
                                )
                            if product_conflict_rows and overwrite:
                                delete_natural_slice = True
                        if registry_conflict and overwrite and str(registry_conflict.get("category_key") or "") != str(category_key):
                            con.execute("DELETE FROM cube_registry WHERE id = ?", [registry_conflict["id"]])
                    if not exists:
                        con.execute(f"CREATE TABLE {quoted_table} AS SELECT * FROM {quote_identifier(stage_table)}")
                        inserted = _stage_count(con, stage_table)
                    else:
                        target_types = _ensure_table_accepts_stage_columns(
                            con,
                            table_name=table_name,
                            quoted_table=quoted_table,
                            stage_table=stage_table,
                            stage_columns=stage_columns,
                        )
                        _backfill_business_row_hash(
                            con,
                            quoted_table=quoted_table,
                            stage_columns=stage_columns,
                            project_name=project_name,
                            year=int(year),
                            month=int(month),
                            marketplace_code=marketplace_code,
                        )
                        if effective_category_name and delete_natural_slice:
                            _delete_product_slice_by_natural_key(
                                con,
                                table_name=table_name,
                                quoted_table=quoted_table,
                                project_name=project_name,
                                year=int(year),
                                month=int(month),
                                marketplace_code=marketplace_code,
                                category_name=effective_category_name,
                            )
                        con.execute(
                            f"""
                            DELETE FROM {quoted_table}
                            WHERE {quote_duckdb_name('__project_name')} = ?
                              AND CAST({quote_duckdb_name('__year')} AS INTEGER) = ?
                              AND CAST({quote_duckdb_name('__month')} AS INTEGER) = ?
                              AND {quote_duckdb_name('__marketplace_code')} = ?
                              AND {quote_duckdb_name('__category_key')} = ?
                            """,
                            [project_name, int(year), int(month), marketplace_code, category_key],
                        )
                        inserted = _stage_insertable_count(
                            con,
                            quoted_table=quoted_table,
                            stage_table=stage_table,
                            business_deduplicate=True,
                        )
                        con.execute(
                            _insert_stage_sql(
                                quoted_table=quoted_table,
                                stage_table=stage_table,
                                columns=stage_columns,
                                target_types=target_types,
                                deduplicate=False,
                                business_deduplicate=True,
                            )
                        )

                    con.execute(
                        """
                        INSERT INTO pipeline_loads (
                            table_name, source_file, load_name, project_name, mode, rows_loaded
                        )
                        VALUES (?, ?, ?, ?, ?, ?)
                        """,
                        [
                            table_name,
                            str(csv_path),
                            load_name or f"smart_pipeline:{run_id}",
                            project_name,
                            "replace" if overwrite else "replace_slice",
                            inserted,
                        ],
                    )
            metrics["rows_affected"] = inserted
        return inserted

    def search_products(
        self,
        *,
        table_name: str,
        query_text: str | None = None,
        project_name: str | None = None,
        run_id: str | None = None,
        marketplace: str | None = None,
        category: str | None = None,
        limit: int = 100,
        offset: int = 0,
    ) -> dict[str, Any]:
        quote_identifier(table_name)
        if not self.table_exists(table_name):
            return {"columns": [], "rows": [], "total": 0, "run_id": run_id}

        columns = self.table_columns(table_name)
        quoted_table = quote_identifier(table_name)
        where: list[str] = []
        params: list[Any] = []

        effective_run_id = run_id
        if effective_run_id is None and "__run_id" in columns:
            effective_run_id = self.latest_successful_run_id(project_name=project_name)
        if effective_run_id and "__run_id" in columns:
            where.append(f"{quote_duckdb_name('__run_id')} = ?")
            params.append(effective_run_id)
        if project_name and "__project_name" in columns:
            where.append(f"{quote_duckdb_name('__project_name')} = ?")
            params.append(project_name)
        if marketplace and "Маркетплейс" in columns:
            where.append(f"{quote_duckdb_name('Маркетплейс')} = ?")
            params.append(marketplace)
        if category and "Категория" in columns:
            where.append(f"{quote_duckdb_name('Категория')} = ?")
            params.append(category)
        if query_text:
            searchable = [column for column in SEARCH_COLUMNS if column in columns]
            if searchable:
                needle = f"%{query_text.lower()}%"
                where.append(
                    "("
                    + " OR ".join(
                        f"lower(CAST({quote_duckdb_name(column)} AS VARCHAR)) LIKE ?" for column in searchable
                    )
                    + ")"
                )
                params.extend([needle] * len(searchable))

        where_sql = " WHERE " + " AND ".join(where) if where else ""
        order_sql = f" ORDER BY {quote_duckdb_name('__imported_at')} DESC" if "__imported_at" in columns else ""
        safe_limit = max(1, min(int(limit), 500))
        safe_offset = max(0, int(offset))

        count_row = self._fetch_one(f"SELECT COUNT(*) AS total FROM {quoted_table}{where_sql}", params, read_only=True)
        rows = self._fetch_records(
            f"SELECT * FROM {quoted_table}{where_sql}{order_sql} LIMIT {safe_limit} OFFSET {safe_offset}",
            params,
            read_only=True,
        )
        return {
            "columns": columns,
            "rows": rows,
            "total": int(count_row["total"]) if count_row else 0,
            "run_id": effective_run_id,
        }

    def refresh_large_category_flags(
        self,
        *,
        project_name: str,
        category_keys: list[str] | None = None,
        slice_limit: int = HEAVY_SLICE_ROWS_LIMIT,
        category_limit: int = HEAVY_CATEGORY_ROWS_LIMIT,
    ) -> list[dict[str, Any]]:
        where = ["project_name = ?"]
        params: list[Any] = [project_name]
        if category_keys:
            clean_keys = sorted({str(key) for key in category_keys if str(key).strip()})
            if clean_keys:
                placeholders = ", ".join("?" for _ in clean_keys)
                where.append(f"category_key IN ({placeholders})")
                params.extend(clean_keys)
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            summaries = con.execute(
                f"""
                SELECT category_key, SUM(rows_count) AS category_rows_count
                FROM cube_registry
                WHERE {" AND ".join(where)}
                GROUP BY category_key
                """,
                params,
            ).fetchall()
            for category_key, total_rows in summaries:
                total = int(total_rows or 0)
                con.execute(
                    """
                    UPDATE cube_registry
                    SET
                        is_heavy = rows_count >= ? OR ? >= ?,
                        data_mode = CASE WHEN rows_count >= ? OR ? >= ? THEN 'heavy' ELSE 'standard' END,
                        heavy_reason = CASE
                            WHEN rows_count >= ? THEN 'Срез содержит ' || CAST(rows_count AS VARCHAR) || ' строк: raw XLSX отключён, используй агрегированные отчёты.'
                            WHEN ? >= ? THEN 'Категория содержит ' || CAST(? AS VARCHAR) || ' строк суммарно: используй агрегированные отчёты вместо raw XLSX.'
                            ELSE NULL
                        END
                    WHERE project_name = ? AND category_key = ?
                    """,
                    [
                        int(slice_limit),
                        total,
                        int(category_limit),
                        int(slice_limit),
                        total,
                        int(category_limit),
                        int(slice_limit),
                        total,
                        int(category_limit),
                        total,
                        project_name,
                        str(category_key),
                    ],
                )
        return self.large_category_summary(project_name=project_name)

    def large_category_summary(self, *, project_name: str) -> list[dict[str, Any]]:
        rows = self._fetch_records(
            """
            SELECT
                project_name,
                category_key,
                MIN(category_name) AS category_name,
                MIN(marketplace) AS marketplace,
                MIN(marketplace_code) AS marketplace_code,
                COUNT(*) AS slices_count,
                SUM(rows_count) AS rows_count,
                MAX(rows_count) AS max_slice_rows,
                BOOL_OR(COALESCE(is_heavy, false)) AS is_heavy,
                MAX(saved_to_db_at) AS latest_saved_at,
                MAX(reports_built_at) AS reports_built_at,
                MAX(heavy_reason) AS heavy_reason
            FROM cube_registry
            WHERE project_name = ?
            GROUP BY project_name, category_key
            ORDER BY is_heavy DESC, rows_count DESC, category_name, marketplace
            """,
            [project_name],
        )
        for row in rows:
            row["available_reports"] = ["category_month", "brand_month", "classification_month", "top_sku"]
        return rows

    def report_options(self, *, table_name: str, project_name: str) -> dict[str, Any]:
        if not self.table_exists(table_name):
            return {
                "categories": [],
                "period_from": None,
                "period_to": None,
                "columns": [],
                "warnings": [f"Таблица {table_name} не найдена."],
            }
        columns = self.table_columns(table_name)
        warnings: list[str] = []
        missing = [column for column in EXPORT_METADATA_COLUMNS if column not in columns]
        if missing:
            warnings.append(
                "Агрегированные отчёты доступны после сохранения данных через smart pipeline: не хватает "
                + ", ".join(missing)
                + "."
            )
            return {"categories": [], "period_from": None, "period_to": None, "columns": columns, "warnings": warnings}
        self.refresh_large_category_flags(project_name=project_name)
        period = self._export_period_from_cube(project_name=project_name)
        min_period = int(period["min_period"]) if period and period.get("min_period") is not None else None
        max_period = int(period["max_period"]) if period and period.get("max_period") is not None else None
        return {
            "categories": self.large_category_summary(project_name=project_name),
            "period_from": _period_index_to_label(min_period) if min_period else None,
            "period_to": _period_index_to_label(max_period) if max_period else None,
            "columns": self.export_visible_columns(table_name=table_name, project_name=project_name),
            "warnings": warnings,
        }

    def count_report_rows(
        self,
        *,
        table_name: str,
        project_name: str,
        report_type: str,
        category_keys: list[str] | None = None,
        period_from_index: int | None = None,
        period_to_index: int | None = None,
    ) -> int:
        query, params, _ = self._report_query_sql(
            table_name=table_name,
            project_name=project_name,
            report_type=report_type,
            category_keys=category_keys,
            period_from_index=period_from_index,
            period_to_index=period_to_index,
        )
        row = self._fetch_one(
            f"SELECT COUNT(*) AS total FROM ({query}) report_rows",
            params,
            read_only=True,
            temp_directory=self._duckdb_temp_directory(),
        )
        return int(row["total"]) if row else 0

    def fetch_report_dataframe(
        self,
        *,
        table_name: str,
        project_name: str,
        report_type: str,
        category_keys: list[str] | None = None,
        period_from_index: int | None = None,
        period_to_index: int | None = None,
        limit: int | None = 500,
        offset: int = 0,
    ) -> pd.DataFrame:
        query, params, _ = self._report_query_sql(
            table_name=table_name,
            project_name=project_name,
            report_type=report_type,
            category_keys=category_keys,
            period_from_index=period_from_index,
            period_to_index=period_to_index,
        )
        if limit is not None:
            query = f"{query} LIMIT ? OFFSET ?"
            params = [*params, max(1, int(limit)), max(0, int(offset))]
        with self._lock, connect(self.settings.db_path, read_only=True, temp_directory=self._duckdb_temp_directory()) as con:
            return con.execute(query, params).fetchdf()

    def _load_excel_extension(self, con: Any) -> None:
        con.execute("INSTALL excel")
        con.execute("LOAD excel")

    def _export_flat_query_with_openpyxl(
        self,
        con: Any,
        *,
        query: str,
        output_path: Path,
        params: dict[str, Any] | list[Any],
        header: bool,
        sheet_name: str,
        row_count: int | None,
        batch_size: int = 50_000,
    ) -> int:
        try:
            from openpyxl import Workbook
        except ModuleNotFoundError as exc:
            raise ImportError("Для fallback XLSX нужен openpyxl. Установи зависимости проекта: pip install -r requirements.txt") from exc

        workbook = Workbook(write_only=True)
        worksheet = workbook.create_sheet(_validate_sheet_name(sheet_name))
        metadata = con.execute(f"SELECT * FROM ({query}) AS flat_export_source LIMIT 0", params)
        columns = [str(item[0]) for item in (metadata.description or [])]
        if header:
            worksheet.append(columns)

        written = 0
        total_rows = int(row_count or 0)
        while True:
            limit_sql, execute_params = _query_params_with_limit(params, limit=batch_size, offset=written)
            result = con.execute(
                f"SELECT * FROM ({query}) AS flat_export_source{limit_sql}",
                execute_params,
            )
            rows = result.fetchall()
            if not rows:
                break
            for row in rows:
                worksheet.append([_xlsx_cell_value(value) for value in row])
            written += len(rows)
            if total_rows and written >= total_rows:
                break

        workbook.save(output_path)
        return written

    def export_flat_query(
        self,
        query: str,
        output_path: Path,
        format: Literal["csv", "xlsx"],
        params: dict[str, Any] | list[Any] | None = None,
        delimiter: str = ";",
        header: bool = True,
        sheet_name: str = "Data",
        csv_decimal_separator: Literal["dot", "comma"] = "comma",
    ) -> ExportResult:
        target = Path(output_path).expanduser()
        target.parent.mkdir(parents=True, exist_ok=True)
        clean_format = str(format).lower()
        if clean_format not in {"csv", "xlsx"}:
            raise ValueError("Формат flat export должен быть csv или xlsx.")
        clean_csv_decimal_separator = str(csv_decimal_separator).lower()
        if clean_csv_decimal_separator not in {"dot", "comma"}:
            raise ValueError("CSV decimal separator должен быть dot или comma.")
        clean_query = _clean_copy_query(query)
        clean_delimiter = _validate_csv_delimiter(delimiter)
        clean_sheet_name = _validate_sheet_name(sheet_name)
        execute_params = _normalize_query_params(params)
        started = time.perf_counter()
        row_count: int | None = None
        status = "success"
        fallback_error: str | None = None
        try:
            with measure_duckdb_operation(
                "export_flat_query",
                {"db_path": self.settings.db_path, "file_path": target, "format": clean_format},
            ) as metrics:
                with self._lock, connect(self.settings.db_path, read_only=True, temp_directory=self._duckdb_temp_directory()) as con:
                    _register_dedup_export_functions(con)
                    if clean_format == "csv":
                        csv_query = clean_query
                        columns: list[str] | None = None
                        if header or clean_csv_decimal_separator == "comma":
                            metadata = con.execute(f"SELECT * FROM ({clean_query}) AS export_source LIMIT 0", execute_params)
                            columns = [str(item[0]) for item in (metadata.description or [])]
                        if clean_csv_decimal_separator == "comma":
                            csv_query = _csv_decimal_comma_query(clean_query, columns or [])
                        if header:
                            columns = columns or []
                            header_prefix = _csv_header_prefix(columns, delimiter=clean_delimiter)
                            options = (
                                f"(FORMAT csv, DELIMITER {sql_literal(clean_delimiter)}, HEADER false, "
                                f"PREFIX {sql_literal(header_prefix)}, SUFFIX '\n')"
                            )
                        else:
                            options = f"(FORMAT csv, DELIMITER {sql_literal(clean_delimiter)}, HEADER false)"
                        rows = con.execute(
                            f"COPY ({csv_query}) TO {sql_literal(str(target))} {options}",
                            execute_params,
                        ).fetchall()
                        row_count = _copy_row_count(rows)
                    else:
                        count_row = con.execute(f"SELECT COUNT(*) FROM ({clean_query}) AS flat_export_source", execute_params).fetchone()
                        row_count = int(count_row[0]) if count_row else 0
                        max_rows = XLSX_MAX_DATA_ROWS_WITH_HEADER if header else 1_048_576
                        if row_count > max_rows:
                            raise ValueError(f"XLSX row limit exceeded ({row_count} rows), use CSV.")
                        try:
                            self._load_excel_extension(con)
                            row = con.execute(
                                f"""
                                COPY ({clean_query}) TO {sql_literal(str(target))}
                                WITH (FORMAT xlsx, HEADER {str(bool(header)).lower()}, SHEET {sql_literal(clean_sheet_name)})
                                """,
                                execute_params,
                            ).fetchone()
                            row_count = _copy_first_row_count(row) or row_count
                        except Exception as exc:
                            fallback_error = f"{type(exc).__name__}: {exc}"
                            status = "fallback"
                            LOGGER.warning("DuckDB XLSX COPY failed, falling back to openpyxl: %s", fallback_error)
                            target.unlink(missing_ok=True)
                            row_count = self._export_flat_query_with_openpyxl(
                                con,
                                query=clean_query,
                                output_path=target,
                                params=execute_params,
                                header=header,
                                sheet_name=clean_sheet_name,
                                row_count=row_count,
                            )
                metrics["rows_exported"] = row_count
                metrics["file_size_bytes"] = target.stat().st_size if target.exists() else 0
        except Exception:
            LOGGER.exception("DuckDB flat export failed: %s", target)
            raise
        duration = time.perf_counter() - started
        return ExportResult(
            output_path=target,
            file_size_bytes=target.stat().st_size if target.exists() else 0,
            duration_seconds=duration,
            row_count=row_count,
            status=status,
            format=clean_format,
            error=fallback_error,
        )

    def export_query_to_csv(
        self,
        query: str,
        output_path: Path,
        params: dict[str, Any] | list[Any] | None = None,
        delimiter: str = ";",
        header: bool = True,
    ) -> ExportResult:
        try:
            return self.export_flat_query(query, output_path, "csv", params=params, delimiter=delimiter, header=header)
        except Exception:
            LOGGER.exception("DuckDB COPY CSV export failed: %s", output_path)
            raise

    def export_query_to_xlsx(
        self,
        query: str,
        output_path: Path,
        params: dict[str, Any] | list[Any] | None = None,
        header: bool = True,
        sheet_name: str = "Data",
    ) -> ExportResult:
        return self.export_flat_query(
            query,
            output_path,
            "xlsx",
            params=params,
            header=header,
            sheet_name=sheet_name,
        )

    def export_report_to_csv(
        self,
        *,
        table_name: str,
        target: str | Path,
        project_name: str,
        report_type: str,
        category_keys: list[str] | None = None,
        period_from_index: int | None = None,
        period_to_index: int | None = None,
        limit: int | None = None,
    ) -> ExportResult:
        query, params, _ = self._report_query_sql(
            table_name=table_name,
            project_name=project_name,
            report_type=report_type,
            category_keys=category_keys,
            period_from_index=period_from_index,
            period_to_index=period_to_index,
        )
        if limit is not None:
            query = f"{query} LIMIT ?"
            params = [*params, max(1, int(limit))]
        return self.export_flat_query(query, Path(target), "csv", params=params, delimiter=";", header=True, sheet_name="Data")

    def export_report_to_xlsx(
        self,
        *,
        table_name: str,
        target: str | Path,
        project_name: str,
        report_type: str,
        category_keys: list[str] | None = None,
        period_from_index: int | None = None,
        period_to_index: int | None = None,
        limit: int | None = None,
    ) -> ExportResult:
        query, params, _ = self._report_query_sql(
            table_name=table_name,
            project_name=project_name,
            report_type=report_type,
            category_keys=category_keys,
            period_from_index=period_from_index,
            period_to_index=period_to_index,
        )
        if limit is not None:
            query = f"{query} LIMIT ?"
            params = [*params, max(1, int(limit))]
        return self.export_flat_query(query, Path(target), "xlsx", params=params, header=True, sheet_name="Data")

    def mark_reports_built(self, *, project_name: str, category_keys: list[str] | None = None) -> None:
        where = ["project_name = ?"]
        params: list[Any] = [project_name]
        if category_keys:
            clean_keys = sorted({str(key) for key in category_keys if str(key).strip()})
            if clean_keys:
                placeholders = ", ".join("?" for _ in clean_keys)
                where.append(f"category_key IN ({placeholders})")
                params.extend(clean_keys)
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute(f"UPDATE cube_registry SET reports_built_at = now() WHERE {' AND '.join(where)}", params)

    def export_options(self, *, table_name: str, project_name: str, dedup_enabled: bool = False) -> dict[str, Any]:
        quote_identifier(table_name)
        if not self.table_exists(table_name):
            return {
                "columns": [],
                "selected_columns": [],
                "categories": [],
                "period_from": None,
                "period_to": None,
                "warnings": [f"Таблица {table_name} не найдена."],
            }

        columns = self.table_columns(table_name)
        visible_columns = self.export_visible_columns(
            table_name=table_name,
            project_name=project_name,
            dedup_enabled=dedup_enabled,
        )
        missing = [column for column in EXPORT_METADATA_COLUMNS if column not in columns]
        warnings: list[str] = []
        if missing:
            warnings.append(
                "В таблице нет metadata-колонок нового workflow: "
                + ", ".join(missing)
                + ". Выгрузка доступна после сохранения данных через smart pipeline."
            )
            return {
                "columns": visible_columns,
                "selected_columns": visible_columns,
                "categories": [],
                "period_from": None,
                "period_to": None,
                "warnings": warnings,
            }

        period = self._export_period_from_cube(project_name=project_name) or self._fetch_one(
            f"""
            SELECT
                MIN(CAST({quote_duckdb_name('__year')} AS INTEGER) * 12 + CAST({quote_duckdb_name('__month')} AS INTEGER)) AS min_period,
                MAX(CAST({quote_duckdb_name('__year')} AS INTEGER) * 12 + CAST({quote_duckdb_name('__month')} AS INTEGER)) AS max_period
            FROM {quote_identifier(table_name)}
            WHERE {quote_duckdb_name('__project_name')} = ?
            """,
            [project_name],
        )
        min_period = int(period["min_period"]) if period and period.get("min_period") is not None else None
        max_period = int(period["max_period"]) if period and period.get("max_period") is not None else None
        cube_categories = self._export_categories_from_cube(project_name=project_name)
        return {
            "columns": visible_columns,
            "selected_columns": visible_columns,
            "categories": cube_categories or self.export_categories(table_name=table_name, project_name=project_name),
            "period_from": _period_index_to_label(min_period) if min_period else None,
            "period_to": _period_index_to_label(max_period) if max_period else None,
            "warnings": warnings,
        }

    def export_visible_columns(
        self,
        *,
        table_name: str,
        project_name: str | None = None,
        dedup_enabled: bool | None = None,
    ) -> list[str]:
        columns = self.table_columns(table_name)
        visible = [column for column in columns if not column.startswith("__")]
        if (
            dedup_enabled is None
            and project_name
            and self._project_has_successful_dedup(project_name=project_name)
        ):
            visible.extend(column for column in DEDUP_EXPORT_COLUMNS if column not in visible)
        return visible

    def export_categories(
        self,
        *,
        table_name: str,
        project_name: str,
        category_keys: list[str] | None = None,
    ) -> list[dict[str, Any]]:
        quote_identifier(table_name)
        columns = self.table_columns(table_name)
        if any(column not in columns for column in ("__project_name", "__marketplace_code", "__category_key")):
            return []

        quoted_table = quote_identifier(table_name)
        category_name_expr = (
            f"MIN(CAST({quote_duckdb_name('Категория')} AS VARCHAR))"
            if "Категория" in columns
            else f"CAST({quote_duckdb_name('__category_key')} AS VARCHAR)"
        )
        marketplace_expr = (
            f"MIN(CAST({quote_duckdb_name('Маркетплейс')} AS VARCHAR))"
            if "Маркетплейс" in columns
            else f"CAST({quote_duckdb_name('__marketplace_code')} AS VARCHAR)"
        )
        where = [f"{quote_duckdb_name('__project_name')} = ?"]
        params: list[Any] = [project_name]
        if category_keys:
            placeholders = ", ".join("?" for _ in category_keys)
            where.append(f"{quote_duckdb_name('__category_key')} IN ({placeholders})")
            params.extend(category_keys)
        return self._fetch_records(
            f"""
            SELECT
                CAST({quote_duckdb_name('__category_key')} AS VARCHAR) AS category_key,
                {category_name_expr} AS category_name,
                CAST({quote_duckdb_name('__marketplace_code')} AS VARCHAR) AS marketplace_code,
                {marketplace_expr} AS marketplace,
                COUNT(*) AS rows_count
            FROM {quoted_table}
            WHERE {" AND ".join(where)}
            GROUP BY {quote_duckdb_name('__category_key')}, {quote_duckdb_name('__marketplace_code')}
            ORDER BY category_name, marketplace
            """,
            params,
            read_only=True,
            temp_directory=self._duckdb_temp_directory(),
        )

    def export_breakdown(
        self,
        *,
        table_name: str,
        project_name: str,
        category_keys: list[str] | None = None,
        period_from_index: int | None = None,
        period_to_index: int | None = None,
        filters: list[dict[str, str]] | None = None,
        excluded_row_hashes: list[str] | None = None,
        dedup_enabled: bool | None = None,
    ) -> list[dict[str, Any]]:
        columns = self._export_columns_with_dedup(
            table_name=table_name,
            project_name=project_name,
            dedup_enabled=dedup_enabled,
        )
        self._require_export_metadata(columns)
        source_sql = self._export_source_sql(
            table_name=table_name,
            project_name=project_name,
            columns=columns,
            dedup_enabled=dedup_enabled,
        )
        where_sql, params = self._export_where_sql(
            columns=columns,
            project_name=project_name,
            category_keys=category_keys,
            period_from_index=period_from_index,
            period_to_index=period_to_index,
            filters=filters,
            excluded_row_hashes=excluded_row_hashes,
        )
        category_name_expr = (
            f"MIN(CAST({quote_duckdb_name('Категория')} AS VARCHAR))"
            if "Категория" in columns
            else f"CAST({quote_duckdb_name('__category_key')} AS VARCHAR)"
        )
        marketplace_expr = (
            f"MIN(CAST({quote_duckdb_name('Маркетплейс')} AS VARCHAR))"
            if "Маркетплейс" in columns
            else f"CAST({quote_duckdb_name('__marketplace_code')} AS VARCHAR)"
        )
        year_expr = f"CAST({quote_duckdb_name('__year')} AS INTEGER)"
        month_expr = f"CAST({quote_duckdb_name('__month')} AS INTEGER)"
        return self._fetch_records(
            f"""
            SELECT
                {year_expr} AS year,
                {month_expr} AS month,
                printf('%04d-%02d', {year_expr}, {month_expr}) AS period,
                CAST({quote_duckdb_name('__category_key')} AS VARCHAR) AS category_key,
                {category_name_expr} AS category_name,
                CAST({quote_duckdb_name('__marketplace_code')} AS VARCHAR) AS marketplace_code,
                {marketplace_expr} AS marketplace,
                COUNT(*) AS rows_count
            FROM ({source_sql}) AS export_source
            {where_sql}
            GROUP BY
                {quote_duckdb_name('__year')},
                {quote_duckdb_name('__month')},
                {quote_duckdb_name('__category_key')},
                {quote_duckdb_name('__marketplace_code')}
            ORDER BY year, month, category_name, marketplace
            """,
            params,
            read_only=True,
            temp_directory=self._duckdb_temp_directory(),
            register_dedup_export_functions=bool(dedup_enabled),
        )

    def count_export_products(
        self,
        *,
        table_name: str,
        project_name: str,
        category_keys: list[str] | None = None,
        period_from_index: int | None = None,
        period_to_index: int | None = None,
        filters: list[dict[str, str]] | None = None,
        excluded_row_hashes: list[str] | None = None,
        dedup_enabled: bool | None = None,
    ) -> int:
        columns = self._export_columns_with_dedup(
            table_name=table_name,
            project_name=project_name,
            dedup_enabled=dedup_enabled,
        )
        self._require_export_metadata(columns)
        source_sql = self._export_source_sql(
            table_name=table_name,
            project_name=project_name,
            columns=columns,
            dedup_enabled=dedup_enabled,
        )
        where_sql, params = self._export_where_sql(
            columns=columns,
            project_name=project_name,
            category_keys=category_keys,
            period_from_index=period_from_index,
            period_to_index=period_to_index,
            filters=filters,
            excluded_row_hashes=excluded_row_hashes,
        )
        row = self._fetch_one(
            f"SELECT COUNT(*) AS total FROM ({source_sql}) AS export_source{where_sql}",
            params,
            read_only=True,
            temp_directory=self._duckdb_temp_directory(),
            register_dedup_export_functions=bool(dedup_enabled),
        )
        return int(row["total"]) if row else 0

    def fetch_export_products_dataframe(
        self,
        *,
        table_name: str,
        project_name: str,
        output_columns: list[str],
        category_keys: list[str] | None = None,
        period_from_index: int | None = None,
        period_to_index: int | None = None,
        filters: list[dict[str, str]] | None = None,
        excluded_row_hashes: list[str] | None = None,
        sort_column: str | None = None,
        sort_direction: str = "asc",
        limit: int = 100,
        offset: int = 0,
        include_row_hash: bool = False,
        default_order: bool = True,
        dedup_enabled: bool | None = None,
    ) -> pd.DataFrame:
        columns = self._export_columns_with_dedup(
            table_name=table_name,
            project_name=project_name,
            dedup_enabled=dedup_enabled,
        )
        self._require_export_metadata(columns)
        selected_columns = self._safe_export_columns(columns, output_columns)
        if include_row_hash and "__row_hash" in columns and "__row_hash" not in selected_columns:
            selected_columns = [*selected_columns, "__row_hash"]
        select_sql = ", ".join(quote_duckdb_name(column) for column in selected_columns)
        source_sql = self._export_source_sql(
            table_name=table_name,
            project_name=project_name,
            columns=columns,
            dedup_enabled=dedup_enabled,
        )
        where_sql, params = self._export_where_sql(
            columns=columns,
            project_name=project_name,
            category_keys=category_keys,
            period_from_index=period_from_index,
            period_to_index=period_to_index,
            filters=filters,
            excluded_row_hashes=excluded_row_hashes,
        )
        order_sql = self._export_order_sql(
            columns=columns,
            sort_column=sort_column,
            sort_direction=sort_direction,
            default_order=default_order,
        )
        with self._lock, connect(self.settings.db_path, read_only=True, temp_directory=self._duckdb_temp_directory()) as con:
            _register_dedup_export_functions(con)
            return con.execute(
                f"""
                SELECT {select_sql}
                FROM ({source_sql}) AS export_source
                {where_sql}
                {order_sql}
                LIMIT ? OFFSET ?
                """,
                [*params, max(1, int(limit)), max(0, int(offset))],
            ).fetchdf()

    def _export_products_query_sql(
        self,
        *,
        table_name: str,
        project_name: str,
        output_columns: list[str],
        category_keys: list[str] | None = None,
        period_from_index: int | None = None,
        period_to_index: int | None = None,
        filters: list[dict[str, str]] | None = None,
        excluded_row_hashes: list[str] | None = None,
        sort_column: str | None = None,
        sort_direction: str = "asc",
        default_order: bool = False,
        limit: int | None = None,
        offset: int = 0,
        dedup_enabled: bool | None = None,
    ) -> tuple[str, list[Any], list[str]]:
        columns = self._export_columns_with_dedup(
            table_name=table_name,
            project_name=project_name,
            dedup_enabled=dedup_enabled,
        )
        self._require_export_metadata(columns)
        selected_columns = self._safe_export_columns(columns, output_columns)
        select_sql = ", ".join(_raw_export_column_expr(column) for column in selected_columns)
        source_sql = self._export_source_sql(
            table_name=table_name,
            project_name=project_name,
            columns=columns,
            dedup_enabled=dedup_enabled,
        )
        where_sql, params = self._export_where_sql(
            columns=columns,
            project_name=project_name,
            category_keys=category_keys,
            period_from_index=period_from_index,
            period_to_index=period_to_index,
            filters=filters,
            excluded_row_hashes=excluded_row_hashes,
        )
        order_sql = self._export_order_sql(
            columns=columns,
            sort_column=sort_column,
            sort_direction=sort_direction,
            default_order=default_order,
        )
        query = f"""
            SELECT {select_sql}
            FROM ({source_sql}) AS export_source
            {where_sql}
            {order_sql}
        """
        if limit is not None:
            query = f"{query} LIMIT ? OFFSET ?"
            params = [*params, max(1, int(limit)), max(0, int(offset))]
        return query, params, selected_columns

    def export_products_flat(
        self,
        *,
        table_name: str,
        target: str | Path,
        format: Literal["csv", "xlsx"],
        project_name: str,
        output_columns: list[str],
        category_keys: list[str] | None = None,
        period_from_index: int | None = None,
        period_to_index: int | None = None,
        filters: list[dict[str, str]] | None = None,
        excluded_row_hashes: list[str] | None = None,
        sort_column: str | None = None,
        sort_direction: str = "asc",
        default_order: bool = False,
        limit: int | None = None,
        offset: int = 0,
        dedup_enabled: bool | None = None,
    ) -> ExportResult:
        query, params, _ = self._export_products_query_sql(
            table_name=table_name,
            project_name=project_name,
            output_columns=output_columns,
            category_keys=category_keys,
            period_from_index=period_from_index,
            period_to_index=period_to_index,
            filters=filters,
            excluded_row_hashes=excluded_row_hashes,
            sort_column=sort_column,
            sort_direction=sort_direction,
            default_order=default_order,
            limit=limit,
            offset=offset,
            dedup_enabled=dedup_enabled,
        )
        return self.export_flat_query(
            query,
            Path(target),
            format,
            params=params,
            delimiter=";",
            header=True,
            sheet_name="Data",
        )

    def export_products_to_csv(
        self,
        *,
        table_name: str,
        target: str | Path,
        project_name: str,
        output_columns: list[str],
        category_keys: list[str] | None = None,
        period_from_index: int | None = None,
        period_to_index: int | None = None,
        filters: list[dict[str, str]] | None = None,
        excluded_row_hashes: list[str] | None = None,
        sort_column: str | None = None,
        sort_direction: str = "asc",
        default_order: bool = False,
        dedup_enabled: bool | None = None,
    ) -> Path:
        result = self.export_products_flat(
            table_name=table_name,
            target=target,
            format="csv",
            project_name=project_name,
            output_columns=output_columns,
            category_keys=category_keys,
            period_from_index=period_from_index,
            period_to_index=period_to_index,
            filters=filters,
            excluded_row_hashes=excluded_row_hashes,
            sort_column=sort_column,
            sort_direction=sort_direction,
            default_order=default_order,
            dedup_enabled=dedup_enabled,
        )
        return result.output_path

    def _report_query_sql(
        self,
        *,
        table_name: str,
        project_name: str,
        report_type: str,
        category_keys: list[str] | None,
        period_from_index: int | None,
        period_to_index: int | None,
    ) -> tuple[str, list[Any], list[str]]:
        columns = self._export_columns_with_dedup(table_name=table_name, project_name=project_name)
        self._require_export_metadata(columns)
        clean_report_type = report_type if report_type in {"category_month", "brand_month", "classification_month", "top_sku"} else "category_month"
        source_sql = self._export_source_sql(table_name=table_name, project_name=project_name, columns=columns)
        where_sql, params = self._export_where_sql(
            columns=columns,
            project_name=project_name,
            category_keys=category_keys,
            period_from_index=period_from_index,
            period_to_index=period_to_index,
            filters=[],
            excluded_row_hashes=[],
        )

        year_expr = f"CAST({quote_duckdb_name('__year')} AS INTEGER)"
        month_expr = f"CAST({quote_duckdb_name('__month')} AS INTEGER)"
        marketplace_expr = self._report_text_expr(columns, "Маркетплейс", fallback_column="__marketplace_code")
        category_expr = self._report_text_expr(columns, "Категория", fallback_column="__category_key")
        dimensions: list[tuple[str, str]] = [
            ("Период", f"printf('%04d-%02d', {year_expr}, {month_expr})"),
            ("Год", year_expr),
            ("Месяц", month_expr),
            ("Маркетплейс", marketplace_expr),
            ("Категория", category_expr),
        ]
        if clean_report_type == "brand_month":
            dimensions.extend(self._report_optional_dimensions(columns, ("Бренд",)))
        elif clean_report_type == "classification_month":
            dimensions.extend(self._report_optional_dimensions(columns, REPORT_CLASSIFICATION_COLUMNS))
        elif clean_report_type == "top_sku":
            dimensions.extend(
                self._report_optional_dimensions(
                    columns,
                    (
                        "SKU",
                        "Название",
                        "Бренд",
                        "Тип",
                        "Подкатегория",
                        "ML-группа товара",
                        "ML-группа фасовки",
                        "ML-канонический SKU",
                        "ML-dedup статус",
                    ),
                )
            )

        sales_expr = self._report_sum_expr(columns, CUBE_SALES_FILTER_COLUMNS)
        revenue_expr = self._report_sum_expr(columns, REPORT_REVENUE_COLUMNS)
        volume_expr = self._report_volume_kg_expr(columns)
        sku_expr = (
            f"COUNT(DISTINCT NULLIF(TRIM(CAST({quote_duckdb_name('SKU')} AS VARCHAR)), ''))"
            if "SKU" in columns
            else "CAST(0 AS BIGINT)"
        )
        select_parts = [f"{expr} AS {quote_duckdb_name(alias)}" for alias, expr in dimensions]
        metric_parts = [
            f"COUNT(*) AS {quote_duckdb_name('Строк')}",
            f"{sku_expr} AS {quote_duckdb_name('Уникальных SKU')}",
            f"{sales_expr} AS {quote_duckdb_name('Продажи, шт')}",
            f"{revenue_expr} AS {quote_duckdb_name('Выручка, руб')}",
            f"{volume_expr} AS {quote_duckdb_name('Объем, кг')}",
            (
                f"CASE WHEN {sales_expr} > 0 THEN {revenue_expr} / NULLIF({sales_expr}, 0) ELSE NULL END "
                f"AS {quote_duckdb_name('Средняя цена, руб')}"
            ),
            (
                f"CASE WHEN {volume_expr} > 0 THEN {revenue_expr} / NULLIF({volume_expr}, 0) ELSE NULL END "
                f"AS {quote_duckdb_name('Цена за кг')}"
            ),
        ]
        group_parts = [expr for _, expr in dimensions]
        group_sql = ", ".join(group_parts)
        order_sql = self._report_order_sql(clean_report_type, [alias for alias, _ in dimensions])
        query = f"""
            SELECT {", ".join(select_parts + metric_parts)}
            FROM ({source_sql}) AS report_source
            {where_sql}
            GROUP BY {group_sql}
            {order_sql}
        """
        return query, params, [alias for alias, _ in dimensions] + [
            "Строк",
            "Уникальных SKU",
            "Продажи, шт",
            "Выручка, руб",
            "Объем, кг",
            "Средняя цена, руб",
            "Цена за кг",
        ]

    def _report_optional_dimensions(self, columns: list[str], requested: tuple[str, ...]) -> list[tuple[str, str]]:
        return [(column, self._report_text_expr(columns, column)) for column in requested if column in columns]

    def _report_text_expr(self, columns: list[str], column: str, *, fallback_column: str | None = None) -> str:
        if column in columns:
            base = f"NULLIF(TRIM(CAST({quote_duckdb_name(column)} AS VARCHAR)), '')"
        else:
            base = "NULL"
        if fallback_column and fallback_column in columns:
            fallback = f"NULLIF(TRIM(CAST({quote_duckdb_name(fallback_column)} AS VARCHAR)), '')"
            return f"COALESCE({base}, {fallback}, 'Не заполнено')"
        return f"COALESCE({base}, 'Не заполнено')"

    def _report_sum_expr(self, columns: list[str], candidates: tuple[str, ...]) -> str:
        column = next((candidate for candidate in candidates if candidate in columns), None)
        if not column:
            return "CAST(0 AS DOUBLE)"
        return f"COALESCE(SUM({self._report_number_expr(column)}), 0)"

    def _report_volume_kg_expr(self, columns: list[str]) -> str:
        kg_column = next((candidate for candidate in REPORT_VOLUME_KG_COLUMNS if candidate in columns), None)
        if kg_column:
            return f"COALESCE(SUM({self._report_number_expr(kg_column)}), 0)"
        ton_column = next((candidate for candidate in REPORT_VOLUME_T_COLUMNS if candidate in columns), None)
        if ton_column:
            return f"COALESCE(SUM({self._report_number_expr(ton_column)} * 1000), 0)"
        return "CAST(0 AS DOUBLE)"

    @staticmethod
    def _report_number_expr(column: str) -> str:
        quoted = quote_duckdb_name(column)
        nbsp = "\u00a0"
        return (
            "TRY_CAST("
            f"REPLACE(REPLACE(REPLACE(CAST({quoted} AS VARCHAR), '{nbsp}', ''), ' ', ''), ',', '.') "
            "AS DOUBLE)"
        )

    @staticmethod
    def _report_order_sql(report_type: str, dimensions: list[str]) -> str:
        if report_type == "top_sku":
            return f"ORDER BY {quote_duckdb_name('Выручка, руб')} DESC NULLS LAST, {quote_duckdb_name('Продажи, шт')} DESC NULLS LAST"
        order_columns = [
            column
            for column in (
                "Год",
                "Месяц",
                "Категория",
                "Маркетплейс",
                "Бренд",
                "Тип",
                "Подкатегория",
                "ML-группа товара",
                "ML-группа фасовки",
            )
            if column in dimensions
        ]
        if not order_columns:
            return ""
        return "ORDER BY " + ", ".join(quote_duckdb_name(column) for column in order_columns)

    def _project_has_successful_dedup(self, *, project_name: str) -> bool:
        if not self.table_exists("dedup_runs"):
            return False
        row = self._fetch_one(
            """
            SELECT COUNT(*) AS runs_count
            FROM dedup_runs
            WHERE project_name = ? AND status = 'success'
            """,
            [project_name],
            read_only=True,
            temp_directory=self._duckdb_temp_directory(),
        )
        return bool(row and int(row.get("runs_count") or 0) > 0)

    def _export_dedup_mode(self, dedup_enabled: bool | None) -> str:
        if dedup_enabled is None:
            return "columns"
        return "family_sku" if dedup_enabled else "off"

    def _export_columns_with_dedup(
        self,
        *,
        table_name: str,
        project_name: str,
        dedup_enabled: bool | None = None,
    ) -> list[str]:
        columns = self.table_columns(table_name)
        if (
            self._export_dedup_mode(dedup_enabled) == "columns"
            and self._project_has_successful_dedup(project_name=project_name)
        ):
            return [*columns, *(column for column in DEDUP_EXPORT_COLUMNS if column not in columns)]
        return columns

    def _dedup_article_expr(self, columns: list[str], *, table_alias: str) -> str:
        pieces: list[str] = []
        if "Артикул" in columns:
            pieces.append(f"NULLIF(TRIM(CAST({table_alias}.{quote_duckdb_name('Артикул')} AS VARCHAR)), '')")
        if "__business_row_hash" in columns:
            pieces.append(f"NULLIF(TRIM(CAST({table_alias}.{quote_duckdb_name('__business_row_hash')} AS VARCHAR)), '')")
        if "__row_hash" in columns:
            pieces.append(f"NULLIF(TRIM(CAST({table_alias}.{quote_duckdb_name('__row_hash')} AS VARCHAR)), '')")
        if not pieces:
            return "NULL"
        return "COALESCE(" + ", ".join(pieces) + ")"

    def _export_base_select_sql(self, columns: list[str], *, sku_replacement_sql: str | None = None) -> str:
        parts: list[str] = []
        for column in columns:
            quoted = quote_duckdb_name(column)
            if sku_replacement_sql and column == "SKU":
                parts.append(f"COALESCE({sku_replacement_sql}, p.{quoted}) AS {quoted}")
            else:
                parts.append(f"p.{quoted} AS {quoted}")
        return ",\n                ".join(parts)

    def _export_source_sql(
        self,
        *,
        table_name: str,
        project_name: str,
        columns: list[str],
        dedup_enabled: bool | None = None,
    ) -> str:
        quoted_table = quote_identifier(table_name)
        mode = self._export_dedup_mode(dedup_enabled)
        if mode == "off" or not self._project_has_successful_dedup(project_name=project_name):
            return f"SELECT * FROM {quoted_table}"

        include_dedup_columns = mode == "columns"
        replace_sku = mode == "family_sku" and "SKU" in self.table_columns(table_name)
        if mode == "family_sku" and not replace_sku:
            return f"SELECT * FROM {quoted_table}"

        table_columns = self.table_columns(table_name)
        base_select = self._export_base_select_sql(
            table_columns,
            sku_replacement_sql="d.family_sku" if replace_sku else None,
        )
        dedup_select = ""
        if include_dedup_columns:
            dedup_select = f""",
                d.ml_family_id AS {quote_duckdb_name('ML-группа товара')},
                d.ml_pack_id AS {quote_duckdb_name('ML-группа фасовки')},
                d.canonical_sku AS {quote_duckdb_name('ML-канонический SKU')},
                d.ml_dedup_status AS {quote_duckdb_name('ML-dedup статус')},
                d.run_id AS {quote_duckdb_name('ML-dedup run')}"""
        family_ctes = ""
        family_join = ""
        family_select = ""
        if replace_sku:
            family_ctes = f""",
            family_candidates AS (
                SELECT
                    latest.run_id,
                    g.ml_family_id,
                    g.canonical_sku,
                    MIN(n.brand) AS brand,
                    MIN(n.category_name) AS category_name,
                    COALESCE(SUM(n.revenue), 0) AS revenue,
                    COALESCE(SUM(n.sales_volume), 0) AS sales_volume,
                    COUNT(*) AS source_count
                FROM latest_dedup_runs AS latest
                JOIN dedup_sku_nodes AS n
                  ON n.run_id = latest.run_id
                JOIN dedup_sku_groups AS g
                  ON g.run_id = n.run_id AND g.node_id = n.node_id
                GROUP BY
                    latest.run_id,
                    g.ml_family_id,
                    g.canonical_sku
            ),
            family_titles AS (
                SELECT run_id, ml_family_id, family_sku
                FROM (
                    SELECT
                        run_id,
                        ml_family_id,
                        NULLIF({DEDUP_EXPORT_FAMILY_TITLE_FUNCTION}(canonical_sku, brand, category_name), '') AS family_sku,
                        ROW_NUMBER() OVER (
                            PARTITION BY run_id, ml_family_id
                            ORDER BY revenue DESC NULLS LAST, sales_volume DESC NULLS LAST, source_count DESC NULLS LAST, canonical_sku NULLS LAST
                        ) AS family_rank
                    FROM family_candidates
                )
                WHERE family_rank = 1
            )"""
            family_join = """
                LEFT JOIN family_titles AS f
                  ON f.run_id = g.run_id AND f.ml_family_id = g.ml_family_id"""
            family_select = """,
                    f.family_sku"""

        article_expr = self._dedup_article_expr(self.table_columns(table_name), table_alias="p")
        return f"""
            WITH latest_dedup_runs AS (
                SELECT run_id, project_name, category_key
                FROM (
                    SELECT
                        run_id,
                        project_name,
                        category_key,
                        ROW_NUMBER() OVER (
                            PARTITION BY project_name, category_key
                            ORDER BY COALESCE(finished_at, created_at) DESC, created_at DESC
                        ) AS rn
                    FROM dedup_runs
                    WHERE status = 'success'
                      AND project_name = {sql_literal(project_name)}
                )
                WHERE rn = 1
            ){family_ctes},
            dedup_map AS (
                SELECT
                    n.project_name,
                    n.category_key,
                    n.marketplace_code,
                    n.article,
                    g.ml_family_id,
                    g.ml_pack_id,
                    g.canonical_sku,
                    g.ml_dedup_status,
                    g.run_id{family_select}
                FROM latest_dedup_runs AS latest
                JOIN dedup_sku_nodes AS n
                  ON n.run_id = latest.run_id
                JOIN dedup_sku_groups AS g
                  ON g.run_id = n.run_id AND g.node_id = n.node_id
                {family_join}
            )
            SELECT
                {base_select}{dedup_select}
            FROM {quoted_table} AS p
            LEFT JOIN dedup_map AS d
              ON d.project_name = p.{quote_duckdb_name('__project_name')}
             AND d.category_key = p.{quote_duckdb_name('__category_key')}
             AND d.marketplace_code = p.{quote_duckdb_name('__marketplace_code')}
             AND d.article = {article_expr}
        """

    def _require_export_metadata(self, columns: list[str]) -> None:
        missing = [column for column in EXPORT_METADATA_COLUMNS if column not in columns]
        if missing:
            raise ValueError(
                "Для выгрузки нужны metadata-колонки нового workflow: "
                + ", ".join(missing)
                + ". Сохрани данные в БД через smart pipeline."
            )

    def _safe_export_columns(self, columns: list[str], requested: list[str] | None) -> list[str]:
        visible = [column for column in columns if not column.startswith("__")]
        selected = [column for column in (requested or visible) if column in visible]
        return selected or visible

    def _export_where_sql(
        self,
        *,
        columns: list[str],
        project_name: str,
        category_keys: list[str] | None,
        period_from_index: int | None,
        period_to_index: int | None,
        filters: list[dict[str, str]] | None,
        excluded_row_hashes: list[str] | None,
    ) -> tuple[str, list[Any]]:
        where = [f"{quote_duckdb_name('__project_name')} = ?"]
        params: list[Any] = [project_name]
        if category_keys:
            placeholders = ", ".join("?" for _ in category_keys)
            where.append(f"{quote_duckdb_name('__category_key')} IN ({placeholders})")
            params.extend(category_keys)
        period_expr = f"CAST({quote_duckdb_name('__year')} AS INTEGER) * 12 + CAST({quote_duckdb_name('__month')} AS INTEGER)"
        if period_from_index is not None:
            where.append(f"{period_expr} >= ?")
            params.append(int(period_from_index))
        if period_to_index is not None:
            where.append(f"{period_expr} <= ?")
            params.append(int(period_to_index))
        for item in filters or []:
            column = str(item.get("column") or "")
            value = str(item.get("value") or "").strip()
            match_type = str(item.get("match_type") or "contains")
            if not value or column not in columns or column.startswith("__"):
                continue
            column_expr = f"lower(CAST({quote_duckdb_name(column)} AS VARCHAR))"
            lowered = value.lower()
            if match_type == "equals":
                where.append(f"CAST({quote_duckdb_name(column)} AS VARCHAR) = ?")
                params.append(value)
            elif match_type == "not_contains":
                where.append(f"{column_expr} NOT LIKE ?")
                params.append(f"%{lowered}%")
            elif match_type == "startswith":
                where.append(f"{column_expr} LIKE ?")
                params.append(f"{lowered}%")
            elif match_type in {"gt", "gte", "lt", "lte"}:
                numeric_expr = f"TRY_CAST(REPLACE(CAST({quote_duckdb_name(column)} AS VARCHAR), ',', '.') AS DOUBLE)"
                op = {"gt": ">", "gte": ">=", "lt": "<", "lte": "<="}[match_type]
                where.append(
                    f"{numeric_expr} IS NOT NULL AND {numeric_expr} {op} "
                    "TRY_CAST(REPLACE(CAST(? AS VARCHAR), ',', '.') AS DOUBLE)"
                )
                params.append(value)
            else:
                where.append(f"{column_expr} LIKE ?")
                params.append(f"%{lowered}%")
        if excluded_row_hashes:
            hashes = [item for item in excluded_row_hashes if item]
            if hashes:
                placeholders = ", ".join("?" for _ in hashes)
                where.append(f"{quote_duckdb_name('__row_hash')} NOT IN ({placeholders})")
                params.extend(hashes)
        return (" WHERE " + " AND ".join(where), params)

    def _export_period_from_cube(self, *, project_name: str) -> dict[str, Any] | None:
        row = self._fetch_one(
            """
            SELECT
                MIN(year * 12 + month) AS min_period,
                MAX(year * 12 + month) AS max_period,
                COUNT(*) AS slices_count
            FROM cube_registry
            WHERE project_name = ?
            """,
            [project_name],
        )
        if not row or not row.get("slices_count"):
            return None
        return row

    def _export_categories_from_cube(self, *, project_name: str, category_keys: list[str] | None = None) -> list[dict[str, Any]]:
        where = ["project_name = ?"]
        params: list[Any] = [project_name]
        if category_keys:
            placeholders = ", ".join("?" for _ in category_keys)
            where.append(f"category_key IN ({placeholders})")
            params.extend(category_keys)
        return self._fetch_records(
            f"""
            SELECT
                CAST(category_key AS VARCHAR) AS category_key,
                MIN(CAST(category_name AS VARCHAR)) AS category_name,
                CAST(marketplace_code AS VARCHAR) AS marketplace_code,
                MIN(CAST(marketplace AS VARCHAR)) AS marketplace,
                SUM(rows_count) AS rows_count
            FROM cube_registry
            WHERE {" AND ".join(where)}
            GROUP BY category_key, marketplace_code
            ORDER BY category_name, marketplace
            """,
            params,
        )

    def _export_order_sql(
        self,
        *,
        columns: list[str],
        sort_column: str | None,
        sort_direction: str,
        default_order: bool = True,
    ) -> str:
        direction = "DESC" if str(sort_direction).lower() == "desc" else "ASC"
        if sort_column and sort_column in columns and not sort_column.startswith("__"):
            return f" ORDER BY {quote_duckdb_name(sort_column)} {direction} NULLS LAST"
        if not default_order:
            return ""
        default_columns = [
            column
            for column in ("__year", "__month", "Категория", "Маркетплейс", "Название", "SKU")
            if column in columns
        ]
        if not default_columns:
            return ""
        return " ORDER BY " + ", ".join(quote_duckdb_name(column) for column in default_columns)

    def list_schedules(self) -> list[dict[str, Any]]:
        return self._fetch_records("SELECT * FROM app_schedules ORDER BY created_at DESC")

    def get_schedule(self, schedule_id: str) -> dict[str, Any] | None:
        return self._fetch_one("SELECT * FROM app_schedules WHERE schedule_id = ?", [schedule_id])

    def create_schedule(
        self,
        *,
        schedule_id: str,
        name: str,
        project_name: str,
        steps: str,
        enabled: bool,
        interval_minutes: int,
        next_run_at: datetime,
        write_xlsx: bool,
        max_weight_kg: float,
        fill_unclassified: dict[str, Any] | None,
    ) -> dict[str, Any]:
        payload = json.dumps(fill_unclassified, ensure_ascii=False) if fill_unclassified else None
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute(
                """
                INSERT INTO app_schedules (
                    schedule_id, name, project_name, steps, enabled, interval_minutes,
                    next_run_at, write_xlsx, max_weight_kg, fill_unclassified_json
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    schedule_id,
                    name,
                    project_name,
                    steps,
                    enabled,
                    interval_minutes,
                    next_run_at,
                    write_xlsx,
                    max_weight_kg,
                    payload,
                ],
            )
        return self.get_schedule(schedule_id) or {}

    def update_schedule(self, schedule_id: str, values: dict[str, Any]) -> dict[str, Any] | None:
        allowed = {
            "name",
            "project_name",
            "steps",
            "enabled",
            "interval_minutes",
            "next_run_at",
            "last_run_at",
            "write_xlsx",
            "max_weight_kg",
            "fill_unclassified_json",
        }
        assignments = [key for key in values if key in allowed]
        if not assignments:
            return self.get_schedule(schedule_id)
        sql = ", ".join(f"{key} = ?" for key in assignments)
        params = [values[key] for key in assignments]
        params.append(schedule_id)
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            con.execute(
                f"UPDATE app_schedules SET {sql}, updated_at = now() WHERE schedule_id = ?",
                params,
            )
        return self.get_schedule(schedule_id)

    def delete_schedule(self, schedule_id: str) -> bool:
        with self._lock, connect(self.settings.db_path) as con:
            apply_migrations(con)
            before = con.execute("SELECT COUNT(*) FROM app_schedules WHERE schedule_id = ?", [schedule_id]).fetchone()[0]
            con.execute("DELETE FROM app_schedules WHERE schedule_id = ?", [schedule_id])
        return bool(before)

    def due_schedules(self, now: datetime) -> list[dict[str, Any]]:
        return self._fetch_records(
            """
            SELECT *
            FROM app_schedules
            WHERE enabled = true AND next_run_at <= ?
            ORDER BY next_run_at ASC
            """,
            [now],
        )

    def mark_schedule_triggered(self, schedule_id: str, *, now: datetime, interval_minutes: int) -> None:
        self.update_schedule(
            schedule_id,
            {
                "last_run_at": now,
                "next_run_at": now + timedelta(minutes=interval_minutes),
            },
        )
